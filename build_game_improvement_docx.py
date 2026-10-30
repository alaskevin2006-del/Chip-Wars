from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "output"
ASSET_DIR = OUT_DIR / "_assets_word"
QA_DIR = OUT_DIR / "_qa_word"
DOCX_PATH = OUT_DIR / "芯片博弈游戏界面与不完全信息机制改进方案.docx"

INK = RGBColor(25, 31, 34)
MUTED = RGBColor(86, 98, 103)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(0, 0, 0)
HEADER_FILL = "213D4A"
ALT_FILL = "F2F6F7"
BORDER_FILL = "D5DEE1"
ACCENT = RGBColor(176, 121, 35)
BODY_FONT = "Microsoft YaHei"
_number_counter = 0


def ensure_dirs() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    QA_DIR.mkdir(parents=True, exist_ok=True)


def pil_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    path = (
        r"C:\Windows\Fonts\msyhbd.ttc"
        if bold
        else r"C:\Windows\Fonts\msyh.ttc"
    )
    return ImageFont.truetype(path, size=size, index=0)


def text_size(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont) -> tuple[int, int]:
    box = draw.multiline_textbbox((0, 0), text, font=font, spacing=8)
    return box[2] - box[0], box[3] - box[1]


def draw_centered_text(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    font: ImageFont.FreeTypeFont,
    fill: str,
    line_spacing: int = 8,
) -> None:
    x1, y1, x2, y2 = box
    width, height = text_size(draw, text, font)
    x = x1 + (x2 - x1 - width) / 2
    y = y1 + (y2 - y1 - height) / 2
    draw.multiline_text(
        (x, y),
        text,
        font=font,
        fill=fill,
        spacing=line_spacing,
        align="center",
    )


def draw_left_text(
    draw: ImageDraw.ImageDraw,
    x: int,
    y: int,
    text: str,
    font: ImageFont.FreeTypeFont,
    fill: str,
    spacing: int = 8,
) -> None:
    draw.multiline_text((x, y), text, font=font, fill=fill, spacing=spacing)


def draw_arrow(
    draw: ImageDraw.ImageDraw,
    start: tuple[int, int],
    end: tuple[int, int],
    fill: str = "#5D7078",
    width: int = 8,
) -> None:
    draw.line((start, end), fill=fill, width=width)
    x1, y1 = start
    x2, y2 = end
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        points = [
            (x2, y2),
            (x2 - direction * 24, y2 - 14),
            (x2 - direction * 24, y2 + 14),
        ]
    else:
        direction = 1 if y2 > y1 else -1
        points = [
            (x2, y2),
            (x2 - 14, y2 - direction * 24),
            (x2 + 14, y2 - direction * 24),
        ]
    draw.polygon(points, fill=fill)


def build_player_loop_figure(path: Path) -> None:
    scale = 2
    width, height = 1600 * scale, 620 * scale
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)

    title_font = pil_font(42 * scale, bold=True)
    step_font = pil_font(34 * scale, bold=True)
    body_font = pil_font(24 * scale)

    draw.text((70 * scale, 45 * scale), "玩家每轮循环", font=title_font, fill="#18262C")
    draw.text(
        (70 * scale, 108 * scale),
        "把不完全信息判断转化为可执行、可复盘的选择",
        font=body_font,
        fill="#5A6B72",
    )

    labels = [
        ("1  看情报", "公开事件、政策信号\n来源与可信度"),
        ("2  押判断", "对手可能类型\n给出本轮概率"),
        ("3  定政策", "3 张政策卡\n兼看公开表态"),
        ("4  看反应", "第三方、企业与\n许可及供应链变化"),
        ("5  读因果", "影响前三、预测对比\n并更新信念"),
    ]

    box_width = 250 * scale
    box_height = 210 * scale
    gap = 42 * scale
    start_x = 78 * scale
    top_y = 230 * scale
    centers: list[tuple[int, int]] = []

    for index, (title, body) in enumerate(labels):
        x1 = start_x + index * (box_width + gap)
        x2 = x1 + box_width
        y1 = top_y
        y2 = y1 + box_height
        fill = "#EAF1F4" if index != 2 else "#F5EBDD"
        outline = "#95AAB3" if index != 2 else "#C89A55"
        draw.rounded_rectangle(
            (x1, y1, x2, y2),
            radius=22 * scale,
            fill=fill,
            outline=outline,
            width=3 * scale,
        )
        draw_centered_text(
            draw,
            (x1 + 14 * scale, y1 + 25 * scale, x2 - 14 * scale, y1 + 82 * scale),
            title,
            step_font,
            "#19262C",
        )
        draw.rectangle(
            (x1 + 22 * scale, y1 + 100 * scale, x2 - 22 * scale, y1 + 102 * scale),
            fill=outline,
        )
        draw_centered_text(
            draw,
            (x1 + 18 * scale, y1 + 118 * scale, x2 - 18 * scale, y2 - 16 * scale),
            body,
            body_font,
            "#4E6067",
        )
        centers.append(((x1 + x2) // 2, (y1 + y2) // 2))

        if index < len(labels) - 1:
            draw_arrow(
                draw,
                (x2 + 6 * scale, (y1 + y2) // 2),
                (x2 + gap - 8 * scale, (y1 + y2) // 2),
            )

    # Feedback loop back to the first step.
    bottom_y = top_y + box_height + 60 * scale
    draw.line(
        (
            centers[-1][0],
            top_y + box_height + 12 * scale,
            centers[-1][0],
            bottom_y,
            centers[0][0],
            bottom_y,
            centers[0][0],
            top_y + box_height + 16 * scale,
        ),
        fill="#B07923",
        width=7 * scale,
        joint="curve",
    )
    draw_arrow(
        draw,
        (centers[0][0], bottom_y - 2 * scale),
        (centers[0][0], top_y + box_height + 12 * scale),
        fill="#B07923",
        width=7 * scale,
    )
    draw.rounded_rectangle(
        (
            625 * scale,
            bottom_y - 32 * scale,
            975 * scale,
            bottom_y + 32 * scale,
        ),
        radius=18 * scale,
        fill="#FFFFFF",
        outline="#B07923",
        width=3 * scale,
    )
    draw_centered_text(
        draw,
        (
            625 * scale,
            bottom_y - 32 * scale,
            975 * scale,
            bottom_y + 32 * scale,
        ),
        "进入下一轮并修正判断",
        body_font,
        "#8B5E19",
    )

    image.save(path, quality=95)


def build_ui_wireframe(path: Path) -> None:
    scale = 2
    width, height = 1560 * scale, 1050 * scale
    image = Image.new("RGB", (width, height), "#F4F7F8")
    draw = ImageDraw.Draw(image)

    font_title = pil_font(34 * scale, bold=True)
    font_subtitle = pil_font(22 * scale, bold=True)
    font_body = pil_font(20 * scale)
    font_small = pil_font(17 * scale)
    font_metric = pil_font(29 * scale, bold=True)

    page = (35 * scale, 28 * scale, 1525 * scale, 1020 * scale)
    draw.rounded_rectangle(page, radius=22 * scale, fill="#FFFFFF", outline="#B8C4C8", width=3 * scale)

    top = (55 * scale, 50 * scale, 1505 * scale, 165 * scale)
    draw.rounded_rectangle(top, radius=14 * scale, fill="#E8EFF1", outline="#C6D1D5", width=2 * scale)
    draw_left_text(draw, 78 * scale, 72 * scale, "第 02 轮", font_title, "#18262C")
    draw_left_text(
        draw,
        78 * scale,
        118 * scale,
        "任务  国产有效交付达到 40，产业压力低于 55",
        font_body,
        "#4E6067",
    )

    metrics = [
        ("可用算力", "49", "+6.5"),
        ("自主度", "62", "B"),
        ("产业压力", "46", "+12"),
    ]
    metric_x = 800 * scale
    for index, (label, value, delta) in enumerate(metrics):
        x1 = metric_x + index * 225 * scale
        x2 = x1 + 205 * scale
        draw.rounded_rectangle(
            (x1, 66 * scale, x2, 147 * scale),
            radius=12 * scale,
            fill="#FFFFFF",
            outline="#C9D4D8",
            width=2 * scale,
        )
        draw_left_text(draw, x1 + 14 * scale, 78 * scale, label, font_small, "#65777E")
        draw_left_text(draw, x1 + 14 * scale, 103 * scale, value, font_metric, "#18262C")
        draw_left_text(draw, x1 + 126 * scale, 108 * scale, delta, font_small, "#9A671F")

    map_panel = (55 * scale, 185 * scale, 1085 * scale, 795 * scale)
    intel_panel = (1105 * scale, 185 * scale, 1505 * scale, 795 * scale)
    for panel in (map_panel, intel_panel):
        draw.rounded_rectangle(
            panel,
            radius=16 * scale,
            fill="#FFFFFF",
            outline="#C6D1D5",
            width=2 * scale,
        )

    draw_left_text(draw, 80 * scale, 208 * scale, "供应链地图", font_title, "#18262C")
    draw_left_text(
        draw,
        80 * scale,
        252 * scale,
        "一次只看一个镜头",
        font_small,
        "#65777E",
    )

    lens_labels = ["安全", "产业", "外交"]
    for index, label in enumerate(lens_labels):
        x1 = 790 * scale + index * 86 * scale
        fill = "#213D4A" if index == 1 else "#E7EEF1"
        text_fill = "#FFFFFF" if index == 1 else "#52656C"
        draw.rounded_rectangle(
            (x1, 215 * scale, x1 + 72 * scale, 255 * scale),
            radius=10 * scale,
            fill=fill,
            outline=fill,
            width=2 * scale,
        )
        draw_centered_text(
            draw,
            (x1, 215 * scale, x1 + 72 * scale, 255 * scale),
            label,
            font_small,
            text_fill,
        )

    nodes = [
        ("芯片设计", "38.9", "稳定"),
        ("工艺良率", "32.2", "承压"),
        ("制造设备", "36.0", "稳定"),
        ("先进封装", "25.0", "瓶颈"),
        ("HBM 配套", "29.0", "偏紧"),
    ]
    node_width = 165 * scale
    node_height = 190 * scale
    node_gap = 28 * scale
    node_y1 = 380 * scale
    for index, (title, value, status) in enumerate(nodes):
        x1 = 82 * scale + index * (node_width + node_gap)
        x2 = x1 + node_width
        y2 = node_y1 + node_height
        fill = "#F6EADD" if index == 3 else "#EEF4F6"
        outline = "#C18A39" if index == 3 else "#A7B9C0"
        draw.rounded_rectangle(
            (x1, node_y1, x2, y2),
            radius=16 * scale,
            fill=fill,
            outline=outline,
            width=3 * scale,
        )
        draw_centered_text(
            draw,
            (x1 + 8 * scale, y1 := node_y1 + 20 * scale, x2 - 8 * scale, y1 + 45 * scale),
            title,
            font_subtitle,
            "#18262C",
        )
        draw_centered_text(
            draw,
            (x1 + 8 * scale, node_y1 + 80 * scale, x2 - 8 * scale, node_y1 + 130 * scale),
            value,
            font_metric,
            "#18262C",
        )
        draw_centered_text(
            draw,
            (x1 + 8 * scale, node_y1 + 145 * scale, x2 - 8 * scale, node_y1 + 178 * scale),
            status,
            font_small,
            "#8B5E19" if index == 3 else "#5F7279",
        )
        if index < len(nodes) - 1:
            draw_arrow(
                draw,
                (x2 + 3 * scale, node_y1 + node_height // 2),
                (x2 + node_gap - 3 * scale, node_y1 + node_height // 2),
                fill="#8DA0A8",
                width=5 * scale,
            )

    draw_left_text(
        draw,
        82 * scale,
        640 * scale,
        "当前瓶颈  先进封装",
        font_subtitle,
        "#8B5E19",
    )
    draw_left_text(
        draw,
        82 * scale,
        682 * scale,
        "点击节点查看库存、替代供给、受影响企业和三回合变化。",
        font_body,
        "#52656C",
    )
    draw_left_text(
        draw,
        82 * scale,
        727 * scale,
        "其余数据默认折叠，避免与主决策争夺注意力。",
        font_body,
        "#52656C",
    )

    draw_left_text(draw, 1130 * scale, 208 * scale, "情报与对手判断", font_title, "#18262C")
    draw_left_text(
        draw,
        1130 * scale,
        264 * scale,
        "公开信号",
        font_subtitle,
        "#18262C",
    )
    draw_left_text(
        draw,
        1130 * scale,
        300 * scale,
        "美国推动日荷同步管制\n来源：盟友磋商  可信度 70%",
        font_body,
        "#52656C",
    )
    draw_left_text(
        draw,
        1130 * scale,
        390 * scale,
        "本轮判断",
        font_subtitle,
        "#18262C",
    )
    draw_left_text(draw, 1130 * scale, 430 * scale, "美国强硬型  65%", font_body, "#5C6F76")
    draw.rounded_rectangle(
        (1130 * scale, 464 * scale, 1468 * scale, 489 * scale),
        radius=12 * scale,
        fill="#E5ECEE",
    )
    draw.rounded_rectangle(
        (1130 * scale, 464 * scale, 1349 * scale, 489 * scale),
        radius=12 * scale,
        fill="#B07923",
    )
    draw_left_text(draw, 1130 * scale, 502 * scale, "务实交换型  35%", font_body, "#5C6F76")
    draw.rounded_rectangle(
        (1130 * scale, 536 * scale, 1468 * scale, 561 * scale),
        radius=12 * scale,
        fill="#E5ECEE",
    )
    draw.rounded_rectangle(
        (1130 * scale, 536 * scale, 1248 * scale, 561 * scale),
        radius=12 * scale,
        fill="#6D8B96",
    )
    draw_left_text(
        draw,
        1130 * scale,
        606 * scale,
        "一条判断题",
        font_subtitle,
        "#18262C",
    )
    draw_left_text(
        draw,
        1130 * scale,
        644 * scale,
        "美国会扩大管制吗？\n本轮押注：70%",
        font_body,
        "#52656C",
    )
    draw_left_text(
        draw,
        1130 * scale,
        718 * scale,
        "结算后显示信念更新\n与预测校准结果。",
        font_small,
        "#7A5A28",
    )

    tray = (55 * scale, 820 * scale, 1505 * scale, 990 * scale)
    draw.rounded_rectangle(
        tray,
        radius=16 * scale,
        fill="#EAF0F2",
        outline="#C4D0D4",
        width=2 * scale,
    )
    draw_left_text(draw, 80 * scale, 843 * scale, "本轮政策手牌", font_title, "#18262C")
    draw_left_text(draw, 80 * scale, 889 * scale, "2 AP  /  3 张主卡  /  更多政策", font_body, "#52656C")

    cards = [
        ("对日市场准入", "外交", "改善协调\n消耗 15 预算"),
        ("收紧镓许可", "管制", "增强姿态\n产业承压"),
        ("国产封装协作", "产业", "直击瓶颈\n消耗 10 预算"),
    ]
    for index, (title, kind, note) in enumerate(cards):
        x1 = 390 * scale + index * 350 * scale
        x2 = x1 + 320 * scale
        draw.rounded_rectangle(
            (x1, 845 * scale, x2, 965 * scale),
            radius=13 * scale,
            fill="#FFFFFF",
            outline="#B7C5CA",
            width=2 * scale,
        )
        draw_left_text(draw, x1 + 16 * scale, 860 * scale, title, font_subtitle, "#18262C")
        draw_left_text(draw, x1 + 16 * scale, 898 * scale, kind, font_small, "#8B5E19")
        draw_left_text(draw, x1 + 105 * scale, 898 * scale, note, font_small, "#52656C")

    image.save(path, quality=95)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=BORDER_FILL, size=4) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_fixed_table_layout(table) -> None:
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_repeat_header_and_layout(table) -> None:
    set_fixed_table_layout(table)
    repeat_table_header(table.rows[0])


def set_cell_text(cell, text: str, bold=False, color=INK, size=9.4) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    run.font.name = BODY_FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    font_size: float = 9.2,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_header_and_layout(table)

    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        set_cell_margins(cell)
        set_cell_shading(cell, HEADER_FILL)
        set_cell_borders(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_text(cell, header, bold=True, color=WHITE, size=font_size)
    prevent_row_split(table.rows[0])

    for row_index, row_data in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for column_index, value in enumerate(row_data):
            cell = row.cells[column_index]
            cell.width = Inches(widths[column_index])
            set_cell_margins(cell)
            set_cell_shading(cell, ALT_FILL if row_index % 2 else "FFFFFF")
            set_cell_borders(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_text(
                cell,
                value,
                bold=(column_index == 0),
                size=font_size,
            )

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def set_style_font(style, size: float, bold=False, color=BLACK) -> None:
    style.font.name = BODY_FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), BODY_FONT)
    r_fonts.set(qn("w:hAnsi"), BODY_FONT)
    r_fonts.set(qn("w:eastAsia"), BODY_FONT)


def clear_theme_color(style, color: RGBColor) -> None:
    r_pr = style.element.get_or_add_rPr()
    color_node = r_pr.find(qn("w:color"))
    if color_node is not None:
        r_pr.remove(color_node)
    style.font.color.rgb = color


def remove_style_borders(style) -> None:
    p_pr = style.element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    set_style_font(normal, 10.8, color=INK)
    normal.paragraph_format.line_spacing = 1.32
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.space_before = Pt(0)

    title = doc.styles["Title"]
    set_style_font(title, 23, bold=True, color=BLACK)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(14)
    title.paragraph_format.line_spacing = 1.05
    remove_style_borders(title)

    subtitle = doc.styles["Subtitle"]
    set_style_font(subtitle, 12.5, bold=False, color=BLACK)
    subtitle.font.italic = False
    clear_theme_color(subtitle, BLACK)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(12)
    subtitle.paragraph_format.line_spacing = 1.15

    heading1 = doc.styles["Heading 1"]
    set_style_font(heading1, 16.5, bold=True, color=BLACK)
    clear_theme_color(heading1, BLACK)
    remove_style_borders(heading1)
    heading1.paragraph_format.space_before = Pt(16)
    heading1.paragraph_format.space_after = Pt(8)
    heading1.paragraph_format.keep_with_next = True
    heading1.paragraph_format.line_spacing = 1.1

    heading2 = doc.styles["Heading 2"]
    set_style_font(heading2, 12.6, bold=True, color=BLACK)
    clear_theme_color(heading2, BLACK)
    remove_style_borders(heading2)
    heading2.paragraph_format.space_before = Pt(11)
    heading2.paragraph_format.space_after = Pt(4)
    heading2.paragraph_format.keep_with_next = True
    heading2.paragraph_format.line_spacing = 1.12

    heading3 = doc.styles["Heading 3"]
    set_style_font(heading3, 11.3, bold=True, color=BLACK)
    clear_theme_color(heading3, BLACK)
    remove_style_borders(heading3)
    heading3.paragraph_format.space_before = Pt(8)
    heading3.paragraph_format.space_after = Pt(3)
    heading3.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        set_style_font(style, 10.6, color=INK)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.76)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.32)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("芯片博弈游戏界面与不完全信息机制改进方案  |  第 ")
    run.font.name = BODY_FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED

    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    field_separate = OxmlElement("w:fldChar")
    field_separate.set(qn("w:fldCharType"), "separate")
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instr_text)
    run._r.append(field_separate)
    run._r.append(field_end)

    suffix = paragraph.add_run(" 页")
    suffix.font.name = BODY_FONT
    suffix._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
    suffix.font.size = Pt(8)
    suffix.font.color.rgb = MUTED


def add_caption(doc: Document, text: str, before: bool = False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(2 if before else 4)
    paragraph.paragraph_format.space_after = Pt(4 if before else 10)
    paragraph.paragraph_format.keep_with_next = before
    run = paragraph.add_run(text)
    run.font.name = BODY_FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(text)


def reset_numbering() -> None:
    global _number_counter
    _number_counter = 0


def add_numbered(doc: Document, text: str) -> None:
    global _number_counter
    _number_counter += 1
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.28)
    paragraph.paragraph_format.first_line_indent = Inches(-0.28)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    paragraph.add_run(f"{_number_counter}.  {text}")


def add_lead_paragraph(doc: Document, lead: str, body: str) -> None:
    paragraph = doc.add_paragraph()
    lead_run = paragraph.add_run(lead)
    lead_run.bold = True
    paragraph.add_run(body)


def add_figure(doc: Document, image_path: Path, width: float, caption: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width))
    add_caption(doc, caption)


def add_page_break(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def build_docx() -> None:
    player_loop = ASSET_DIR / "player_loop.png"
    ui_wireframe = ASSET_DIR / "ui_wireframe.png"
    build_player_loop_figure(player_loop)
    build_ui_wireframe(ui_wireframe)

    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    add_page_number(doc.sections[0].footer.paragraphs[0])

    props = doc.core_properties
    props.title = "芯片博弈游戏界面与不完全信息机制改进方案"
    props.subject = "游戏界面、数据呈现、不完全信息动态博弈与可玩性改版建议"
    props.author = ""
    props.keywords = "芯片博弈, 不完全信息, 动态博弈, 游戏设计, 界面设计"

    # Page 1
    title = doc.add_paragraph(style="Title")
    title_run = title.add_run("芯片博弈游戏界面与")
    title_run.add_break()
    title_run.add_text("不完全信息机制改进方案")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("从模型演示走向可参与的动态政策博弈")

    meta = doc.add_paragraph()
    meta.paragraph_format.space_before = Pt(6)
    meta.paragraph_format.space_after = Pt(16)
    meta_run = meta.add_run("产品改版建议书  |  2026 年 10 月 27 日")
    meta_run.font.name = BODY_FONT
    meta_run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
    meta_run.font.size = Pt(9.5)
    meta_run.font.color.rgb = MUTED

    doc.add_heading("执行摘要", level=1)
    doc.add_paragraph(
        "这款游戏已经具备较完整的政策、企业、第三方政府与供应链结算结构，"
        "也实现了私人类型、分阶段行动和逐回合因果记录。当前的主要矛盾是："
        "模型提供的信息量超过了玩家在单轮决策中能够处理的范围，而支持不完全信息"
        "推理的信息又没有以可操作的方式呈现。"
    )
    doc.add_paragraph(
        "建议保留底层规则复杂度，重构信息层级，并把完美贝叶斯均衡的关键直觉翻译成"
        "玩家每轮都能完成的判断、下注、行动和复盘循环。界面可以更简单，但战略取舍"
        "不应减少。"
    )
    add_lead_paragraph(
        doc,
        "核心改版方向：",
        "用棋盘承载供应链，用情报承载不确定性，用政策卡承载选择，用因果战报承载反馈。",
    )

    doc.add_heading("四项优先动作", level=2)
    reset_numbering()
    add_numbered(doc, "建立三层信息架构。主屏只保留任务、3 项角色指标、3 张政策卡和行动资源，其余数据按需展开。")
    add_numbered(doc, "补齐数据变化表达。每项指标同时显示上一轮值、环比、三回合趋势、目标值和目标差距。")
    add_numbered(doc, "把战报压缩为影响前三。先说明玩家做了什么、哪条供应链发生变化、下一轮要注意什么，完整记录默认折叠。")
    add_numbered(doc, "增加一轮一次的可玩推理。公开信号、对手类型概率、玩家预测、可信度和分支预测必须共同出现。")

    doc.add_paragraph(
        "建议先完成界面可读性和战报重构，再接入信号与信念系统，最后扩展地图、外交、"
        "科技树和多结局。直接重做完整地图并非首选，因为当前最需要验证的是玩家能否"
        "理解并享受一轮博弈。"
    )

    add_page_break(doc)

    # Page 2
    doc.add_heading("一  现状诊断", level=1)
    doc.add_paragraph(
        "现有规则引擎和供应链因果链是项目的核心资产。问题主要集中在信息调度："
        "系统把大量真实变化同时呈现给玩家，却没有把变化转化为可执行的决策线索。"
    )
    doc.add_heading("已有优势", level=2)
    add_bullet(doc, "供应链被拆分为设计、工艺、设备、封装和 HBM 五个约束节点，具备可视化基础。")
    add_bullet(doc, "政府、第三方政府和企业按时序独立决策，能够支撑动态博弈结构。")
    add_bullet(doc, "每轮都有完整的因果记录，便于生成解释、复盘和教学材料。")
    add_bullet(doc, "政策行动同时影响安全、产业、财政和外交，具备真实取舍。")

    doc.add_heading("主要问题", level=2)
    add_caption(doc, "表 1  当前问题与设计回应", before=True)
    add_table(
        doc,
        ["问题", "实际表现", "玩家影响", "设计回应"],
        [
            [
                "信息过载",
                "单屏同时出现 4 项指标、5 个瓶颈、3 个第三方主体、2 类材料和 7 项政策。",
                "无法判断当前最重要的决策。",
                "三层信息架构，主屏只保留任务、3 项指标和 3 张主卡。",
            ],
            [
                "趋势缺失",
                "实测第 2 轮仍显示“首轮基线·暂无趋势”。",
                "无法判断政策是否产生效果。",
                "显示环比、三回合趋势、目标和目标差距。",
            ],
            [
                "战报过长",
                "首轮战报出现 20 条因果记录和 6 个结算模块。",
                "影响最大的事件被淹没。",
                "影响前三、单条因果链、下轮预警，其余折叠。",
            ],
            [
                "研究化文案",
                "界面出现 M1_LEGAL_TRADE、S1_MAJOR_GOV 和 [G] 等标记。",
                "打断沉浸，也增加理解成本。",
                "玩家界面使用自然语言，开发标签放入调试模式。",
            ],
            [
                "推理不可操作",
                "已有隐藏类型与分阶段结算，但缺少公开信号和玩家预测。",
                "不完全信息更像盲选。",
                "信号、信念、预测、可信度和分支预测。",
            ],
            [
                "胜负不清晰",
                "综合绩效分数缺少可解释目标和维度。",
                "玩家不知道战略是否成功。",
                "角色任务、维度评级、阶段目标和多结局。",
            ],
        ],
        [1.02, 2.15, 1.42, 2.18],
        font_size=8.7,
    )

    doc.add_paragraph(
        "因此，真正需要减少的是同一时刻出现的决策对象，而不是模型变量。"
        "只要玩家能在需要时展开数据和因果，复杂度可以保留。"
    )

    add_page_break(doc)

    # Page 3
    doc.add_heading("二  设计原则与玩家循环", level=1)
    doc.add_heading("三项原则", level=2)
    add_lead_paragraph(
        doc,
        "简单不等于战略减少。",
        "界面简化的目标，是让玩家在行动前更容易理解当前局势和真实取舍，"
        "而不是把复杂博弈变成唯一正解。",
    )
    add_lead_paragraph(
        doc,
        "数据必须解释变化。",
        "每项数据都要回答现在怎样、比上轮变了多少、为什么变化，以及距目标还有多远。",
    )
    add_lead_paragraph(
        doc,
        "先呈现不确定性，再呈现结果。",
        "不完全信息博弈的乐趣来自判断和校准，不应只在结算后突然揭示所有对手行动。",
    )

    add_figure(
        doc,
        player_loop,
        6.45,
        "图 1  玩家每轮循环示意图",
    )
    doc.add_paragraph(
        "每轮都应形成一次完整闭环。玩家先读取公开信号，再给出对手类型的概率判断，"
        "选择政策与公开表态，观察第三方和企业反应，最后通过战报修正信念。"
        "当这条循环持续六轮，不完全信息才会从后台模型变成可感知的游戏机制。"
    )
    doc.add_heading("回合节奏", level=2)
    add_bullet(doc, "每轮只突出一个主要事件，避免多个同等重要的信号同时争夺注意力。")
    add_bullet(doc, "每轮只要求一条概率判断，把推理变成短动作，而不是重复填表。")
    add_bullet(doc, "行动选择控制在 3 张主卡，允许玩家在熟悉后再展开更多政策。")

    add_page_break(doc)
    doc.add_heading("信息分层", level=2)
    add_caption(doc, "表 2  游戏信息展示层级", before=True)
    add_table(
        doc,
        ["层级", "内容", "展示方式", "目的"],
        [
            [
                "决策必需",
                "当前任务、轮次、3 项角色指标、AP、预算、3 张主卡。",
                "始终可见。",
                "让玩家快速理解当前局势。",
            ],
            [
                "当前焦点",
                "供应链地图、一个镜头、一个瓶颈、一条事件。",
                "主屏中央或右侧。",
                "把注意力集中到本轮关键矛盾。",
            ],
            [
                "按需展开",
                "完整产业数据、企业利润、对手详情、历史因果。",
                "点击节点、抽屉或复盘模式。",
                "支持深度分析与重玩。",
            ],
            [
                "开发调试",
                "变量名、精确权重、随机种子、公式和日志编号。",
                "开发者模式或导出文件。",
                "保留可检验性，不干扰玩家。",
            ],
        ],
        [1.05, 2.25, 1.55, 1.9],
        font_size=8.9,
    )
    doc.add_paragraph(
        "默认层级应由本轮任务决定。只有当某个数据会改变当前可执行政策时，它才进入主屏；"
        "其余数据保留在节点详情、战报抽屉或复盘界面中。"
    )
    doc.add_heading("按需展开的交互入口", level=2)
    add_table(
        doc,
        ["点击对象", "展开内容"],
        [
            [
                "顶部指标",
                "上轮值、三回合趋势、主要驱动因素、阶段目标与差距。",
            ],
            [
                "供应链节点",
                "库存、替代供给、受影响企业、历史变化与相关许可。",
            ],
            [
                "第三方主体",
                "政策倾向、近期信号、经济依赖、外交筹码和可信度。",
            ],
            [
                "战报条目",
                "完整因果链、相关变量、对企业和第三方的后续影响。",
            ],
        ],
        [1.35, 5.45],
        font_size=9.0,
    )

    add_page_break(doc)

    # Page 4
    doc.add_heading("三  决策主界面", level=1)
    doc.add_paragraph(
        "主界面应从“数据仪表盘”转为“可操作的供应链棋盘”。玩家不需要同时看到所有数字，"
        "只需要看到当前任务、三个关键状态、一个供应链瓶颈和可执行的少量政策。"
    )
    add_figure(
        doc,
        ui_wireframe,
        6.5,
        "图 2  决策主界面布局示意",
    )

    doc.add_heading("页面布局", level=2)
    reset_numbering()
    add_numbered(doc, "顶部显示轮次、阶段任务和 3 项角色专属指标。预算与 AP 固定在底部行动栏。")
    add_numbered(doc, "中部放置供应链地图，支持安全、产业、外交三种镜头，但一次只显示一种。")
    add_numbered(doc, "右侧固定为情报与对手判断，显示公开信号、可信度、类型概率和一条预测题。")
    add_numbered(doc, "底部显示 3 张主政策卡，其他行动放入“更多政策”抽屉。")
    doc.add_heading("主界面禁区", level=2)
    add_bullet(doc, "不在主屏同时展开超过一个警告，否则玩家无法判断优先级。")
    add_bullet(doc, "不让多个指标使用同等醒目的颜色，只突出当前瓶颈和直接变化。")
    add_bullet(doc, "不把完整结算日志直接放在行动界面，战报应在回合结束后单独呈现。")

    role_heading = doc.add_heading("角色专属指标", level=2)
    role_heading.paragraph_format.page_break_before = True
    add_lead_paragraph(
        doc,
        "中国政府：",
        "可用算力、国产有效交付、产业压力。分别回答算力是否够用、国产替代是否兑现、产业是否失血。",
    )
    add_lead_paragraph(
        doc,
        "美国政府：",
        "技术领先、企业损失、盟友协调。分别回答管制是否有效、本国产业是否承压、联盟是否稳固。",
    )
    add_lead_paragraph(
        doc,
        "第三方政府：",
        "安全依赖、市场收益、战略自主。用于表现小国在大国竞争中的真实约束。",
    )
    doc.add_paragraph(
        "每项主指标只保留一个数字和一条变化信息。详细来源、企业和计算公式在点击后显示，"
        "避免主屏再次变成指标墙。"
    )

    doc.add_heading("地图镜头", level=2)
    add_bullet(doc, "安全镜头显示许可、联盟、执法和管制范围。")
    add_bullet(doc, "产业镜头显示产能、瓶颈、企业利润和库存。")
    add_bullet(doc, "外交镜头显示荷兰、日本、韩国以及台积电等关键主体的政策倾向。")

    doc.add_heading("信息优先级检查", level=2)
    add_caption(doc, "表 3  主界面信息优先级", before=True)
    add_table(
        doc,
        ["优先级", "应显示", "不应默认显示"],
        [
            [
                "必须显示",
                "轮次、阶段任务、3 项角色指标、当前瓶颈、3 张主卡、AP 和预算。",
                "完整企业利润表、所有来源说明、研究型变量名。",
            ],
            [
                "按需显示",
                "节点库存、替代供给、历史趋势、第三方政策细节。",
                "跨页日志、所有结算模块、公式和内部编号。",
            ],
            [
                "移动端调整",
                "主指标改为 2 列卡片，主卡横向滑动，底栏固定但不遮挡正文。",
                "多列地图与右侧情报栏同时展开。",
            ],
        ],
        [1.3, 3.15, 2.35],
        font_size=8.9,
    )
    doc.add_heading("视觉规则", level=2)
    add_bullet(doc, "每个屏幕只保留一个主要视觉焦点，其余模块降低对比度。")
    add_bullet(doc, "红色只用于风险上升，绿色只用于状态改善，避免同一颜色表达多种含义。")
    add_bullet(doc, "图标用于识别节点和动作，关键结果仍使用文字和数字表达。")

    add_page_break(doc)

    # Page 5
    doc.add_heading("四  数据变化与战报表达", level=1)
    doc.add_paragraph(
        "玩家不需要每次都读完整模型，但必须能看出“上一轮的选择是否改变了局势”。"
        "指标、政策卡和战报应围绕同一组关键问题设计。"
    )

    doc.add_heading("指标呈现", level=2)
    add_bullet(doc, "当前值与上一轮值并列，变化使用箭头和正负值。")
    add_bullet(doc, "显示三回合迷你趋势，区分短期波动和持续恶化。")
    add_bullet(doc, "显示阶段目标与差距，例如“国产交付 25，目标 40，还差 15”。")
    add_bullet(doc, "点击指标后展开分解原因，不再把所有解释塞进主卡。")

    doc.add_heading("战报结构", level=2)
    reset_numbering()
    add_numbered(doc, "本轮最重要的三个变化，按影响排序，不按模块顺序排列。")
    add_numbered(doc, "玩家政策与变化之间的因果链，突出玩家行动的实际贡献。")
    add_numbered(doc, "对手类型信念是否更新，哪些公开信号改变了判断。")
    add_numbered(doc, "下轮预警，只提示一个新的主要威胁或政策窗口。")
    doc.add_paragraph(
        "完整的 20 条记录可以保留，但默认折叠为“查看完整因果记录”。"
        "玩家需要的是可解释的摘要，而不是把结算日志直接当成叙事。"
    )

    doc.add_heading("文案改写", level=2)
    add_caption(doc, "表 4  研究型界面文案的玩家化改写", before=True)
    add_table(
        doc,
        ["当前表达", "建议表达", "改动目的"],
        [
            [
                "首轮基线·暂无趋势",
                "49；较上轮 +6.5；目标 55",
                "把静态值改成可判断的变化。",
            ],
            [
                "20 条因果记录",
                "本轮 3 个关键变化",
                "先给出优先级，再提供完整记录。",
            ],
            [
                "M1 · 合法贸易",
                "合法芯片交付如何影响企业",
                "去掉内部模块编号。",
            ],
            [
                "综合绩效 61.7",
                "自主度 B；产业压力偏高；距目标还差 6 点交付",
                "让分数对应可理解的战略结果。",
            ],
            [
                "规则决策",
                "对手采取防御性政策；依据其产业压力",
                "把行为翻译成可推断的意图。",
            ],
        ],
        [2.03, 2.8, 2.0],
        font_size=9.0,
    )

    doc.add_heading("政策卡与分支预测", level=2)
    doc.add_paragraph(
        "主政策卡应控制在 3 张以内，每张卡直接列出目标、AP 和预算成本、主要受益方、"
        "主要受损方以及在不同对手类型下的结果区间。玩家选择的是策略组合，"
        "而不是在长列表里寻找按钮。"
    )
    add_caption(doc, "表 5  政策卡的信息结构", before=True)
    add_table(
        doc,
        ["字段", "作用", "示例"],
        [
            [
                "行动目标",
                "让玩家知道这张卡解决什么问题。",
                "缓解先进封装瓶颈。",
            ],
            [
                "资源成本",
                "显示 AP、预算和机会成本。",
                "1 AP，10 预算，占用一个政策槽。",
            ],
            [
                "受益与受损",
                "直接说明谁获益、谁承压。",
                "封装企业获益；财政与材料部门承压。",
            ],
            [
                "信号效果",
                "说明行动会向对手释放什么信息。",
                "被视为产业自救，而非强硬反制。",
            ],
            [
                "分支预测",
                "显示不同对手类型下的结果区间。",
                "对手强硬：交付 31 至 34；对手务实：36 至 39。",
            ],
        ],
        [1.25, 2.6, 2.95],
        font_size=8.9,
    )
    doc.add_heading("动作预览", level=2)
    doc.add_paragraph(
        "预览不应给出唯一结果。系统可以根据当前信念显示两种或三种主要分支，"
        "同时标注可能影响分支的公开信号。玩家据此选择稳健策略，而不是寻找唯一最优按钮。"
    )
    add_bullet(doc, "明确显示结果区间，避免把不确定结果伪装成精确预测。")
    add_bullet(doc, "标出哪一项结果受对手类型影响，哪一项由本国政策直接决定。")
    add_bullet(doc, "结算后同时显示预测区间与实际结果，帮助玩家校准判断。")
    doc.add_heading("示例手牌", level=2)
    add_table(
        doc,
        ["政策卡", "主要收益", "主要代价", "信号效果"],
        [
            [
                "对日市场准入",
                "提高日本维持对冲政策的概率。",
                "消耗预算，国内强硬派压力上升。",
                "被视为合作型信号。",
            ],
            [
                "收紧镓许可",
                "增强反制姿态，获得谈判筹码。",
                "下游成本上升，可能触发对方反制。",
                "被视为高成本强硬信号。",
            ],
            [
                "国产封装协作",
                "直接缓解先进封装瓶颈。",
                "占用预算，短期不改变外交关系。",
                "被视为产业自救信号。",
            ],
        ],
        [1.3, 2.0, 2.15, 1.35],
        font_size=8.7,
    )

    add_page_break(doc)

    # Page 6
    doc.add_heading("五  不完全信息动态博弈的玩家化", level=1)
    doc.add_paragraph(
        "不完全信息动态博弈由三个要素构成：参与者拥有私人类型，通过可观察信号传递或隐藏意图，"
        "其他参与者在行动中更新信念。游戏不需要向玩家展示公式，但必须让这三个要素实际影响决策。"
    )

    doc.add_heading("每轮流程", level=2)
    add_caption(doc, "表 6  可玩的不完全信息回合结构", before=True)
    add_table(
        doc,
        ["阶段", "玩家看到", "玩家执行", "系统处理"],
        [
            [
                "1  公开信号",
                "政策讲话、许可证传闻、盟友磋商、企业游说，以及来源和可信度。",
                "阅读信号，区分廉价表态和昂贵行动。",
                "生成信号，保留部分噪声和误导。",
            ],
            [
                "2  判断下注",
                "对手可能类型及其当前概率。",
                "给出一条预测，例如“美国会扩围 70%”。",
                "记录预测，供赛后校准。",
            ],
            [
                "3  政策选择",
                "3 张主卡、公开表态选项和分支结果区间。",
                "选择政策组合，决定是否释放强硬或合作信号。",
                "检查行动点、预算和条件。",
            ],
            [
                "4  各方反应",
                "第三方企业、盟友政府和本国企业的即时反馈。",
                "观察自己在政策上的公开承诺如何被理解。",
                "按隐藏类型和信念选择反应。",
            ],
            [
                "5  揭示结算",
                "对手实际政策、许可调整和供应链结果。",
                "比较预测与实际。",
                "执行规则结算，更新公开状态。",
            ],
            [
                "6  信念复盘",
                "信念变化、预测校准、关键因果链。",
                "决定下一轮继续坚持还是修正判断。",
                "使用贝叶斯规则更新后验，但不展示公式。",
            ],
        ],
        [1.12, 2.13, 1.8, 1.75],
        font_size=8.6,
    )

    doc.add_heading("隐藏类型设计", level=2)
    add_caption(doc, "表 7  主体隐藏类型与可观察信号", before=True)
    add_table(
        doc,
        ["主体", "建议类型", "可观察信号"],
        [
            [
                "美国政府",
                "安全优先型、产业优先型、条件交换型",
                "盟友协调、许可变化、执法预算、对本国企业损失的容忍度。",
            ],
            [
                "中国政府",
                "反制优先型、发展优先型、条件谈判型",
                "材料许可、产业补贴、市场准入、对外表态的强度。",
            ],
            [
                "荷兰政府",
                "联盟跟随型、市场维护型、战略自主型",
                "DUV 许可、企业豁免、公开政策措辞。",
            ],
            [
                "日本和韩国",
                "安全依赖型、市场依赖型、自主平衡型",
                "设备与 HBM 政策、在华经营安排、同盟会议。",
            ],
        ],
        [1.25, 2.8, 2.75],
        font_size=8.9,
    )
    doc.add_paragraph(
        "隐藏类型必须在整局中保持稳定，只通过信号和行为暴露。若对手策略完全随机，"
        "玩家无法学习，也无法形成可检验的信念。语言模型可以负责生成措辞和事件叙事，"
        "但不应代替类型规则决定核心政策。"
    )

    doc.add_heading("可信度与虚张声势", level=2)
    doc.add_paragraph(
        "公开表态应与实际政策分离。低成本表态允许玩家试探、虚张声势或释放善意，"
        "但反复违背承诺会降低可信度，影响第三方政府和企业的后续判断。"
        "高成本行动，例如启动许可审查或实际部署补贴，则构成更可信的信号。"
    )
    add_caption(doc, "表 8  信号成本与可信度设计", before=True)
    add_table(
        doc,
        ["信号类型", "成本", "玩家用途", "系统反馈"],
        [
            [
                "公开表态",
                "低",
                "试探对手、释放善意、制造模糊。",
                "可信度小幅变化；与实际政策不一致时扣减。",
            ],
            [
                "政策审查与许可",
                "中",
                "表明管制决心，同时保留执行弹性。",
                "更新对手强硬型概率；企业开始调整预期。",
            ],
            [
                "补贴与产能投资",
                "高",
                "展示长期承诺，争取企业与第三方信任。",
                "提高可信度；预算和机会成本立即发生。",
            ],
            [
                "情报与泄露",
                "不确定",
                "提供新证据，迫使玩家判断来源是否可靠。",
                "按来源可信度更新后验，可能包含误导。",
            ],
        ],
        [1.45, 0.95, 2.55, 2.65],
        font_size=8.8,
    )
    doc.add_heading("信念面板", level=2)
    add_bullet(doc, "对手类型：显示强硬、务实或自主三类概率，不用公式。")
    add_bullet(doc, "证据列表：保留最近三条公开信号及其可信度。")
    add_bullet(doc, "预测记录：显示本轮押注、结算结果和误差。")
    add_bullet(doc, "可信度提示：标出玩家自己的表态与实际行动是否一致。")
    doc.add_heading("信念更新示例", level=2)
    add_table(
        doc,
        ["信息", "来源与可信度", "强硬型概率", "界面解释"],
        [
            [
                "开局先验",
                "情景设定",
                "60%",
                "没有额外情报时的初始判断。",
            ],
            [
                "盟友磋商信号",
                "官方渠道，70%",
                "72%",
                "信号支持扩围，但不足以完全确认。",
            ],
            [
                "企业获得临时豁免",
                "企业公告，85%",
                "64%",
                "执行留出空间，判断向务实型回调。",
            ],
        ],
        [1.9, 1.75, 1.2, 1.95],
        font_size=8.8,
    )

    # Page 7
    example_heading = doc.add_heading("六  一轮推理示例与复盘", level=1)
    example_heading.paragraph_format.page_break_before = True
    doc.add_paragraph(
        "以下数字仅用于说明交互结构，不作为真实产业参数。示例展示玩家如何在不看公式的情况下"
        "完成一次不完全信息判断。"
    )

    doc.add_heading("第 2 轮示例", level=2)
    doc.add_paragraph(
        "公开信号显示，美国正在推动日本和荷兰同步扩大设备管制，来源为盟友磋商，可信度 70%。"
        "界面初始判断为美国强硬型 60%、条件交换型 40%；信号出现后更新为强硬型 72%、"
        "条件交换型 28%。玩家需要押注“美国本轮会扩大管制吗”，并给出概率。"
    )
    doc.add_paragraph(
        "政策手牌可包括：对日市场准入承诺、收紧镓许可、国产封装协作。"
        "第一项争取外交空间但消耗预算；第二项释放强硬信号但抬升产业压力；"
        "第三项直接缓解瓶颈，却不改变对手判断。两项行动共同消耗 2 AP，形成真实组合取舍。"
    )
    doc.add_paragraph(
        "行动前预览不给出唯一结果，而是显示两个分支：若日本跟随，国产有效交付预计为 31 至 34，"
        "供应链风险为 43 至 47；若日本保持对冲，交付预计为 36 至 39，风险为 35 至 39。"
        "玩家应选择在不同分支下都可接受的稳健策略。"
    )

    doc.add_heading("结算复盘", level=2)
    add_bullet(doc, "判断是否准确：预测概率与实际结果相差多少。")
    add_bullet(doc, "信念是否合理：哪些证据改变了判断，哪些证据其实不可靠。")
    add_bullet(doc, "政策是否稳健：如果对手选择另一分支，结果是否会失控。")
    add_bullet(doc, "信号是否有效：玩家的公开表态是否改变了第三方的行为。")
    add_bullet(doc, "下一轮是否修正：继续坚持判断，还是转向另一套政策组合。")

    doc.add_paragraph(
        "复盘应同时评价判断质量和结果质量。玩家可能判断正确但结果不佳，也可能因为运气获得高分。"
        "只评价最终结果，会鼓励事后归因和单一最优解，削弱不完全信息博弈的可玩性。"
    )
    doc.add_heading("判断质量与结果质量", level=2)
    add_table(
        doc,
        ["状态", "复盘解释", "下一轮建议"],
        [
            [
                "判断正确，结果良好",
                "证据使用与政策选择一致，但需排除运气成分。",
                "保持判断，同时检查另一分支是否仍可承受。",
            ],
            [
                "判断正确，结果不佳",
                "判断方向正确，但政策组合或资源投入不足。",
                "保留信念，调整政策组合和执行节奏。",
            ],
            [
                "判断错误，结果良好",
                "结果可能来自运气或对手失误。",
                "不要强化错误证据，重新检查信号可信度。",
            ],
            [
                "判断错误，结果不佳",
                "信念和政策同时失配。",
                "降低预测自信，转向更稳健的政策。",
            ],
        ],
        [1.6, 3.1, 2.1],
        font_size=8.8,
    )

    add_page_break(doc)
    doc.add_heading("文明6机制的可借鉴部分", level=2)
    add_caption(doc, "表 9  文明6机制在芯片博弈中的对应设计", before=True)
    add_table(
        doc,
        ["机制", "在本游戏中的用法", "边界"],
        [
            [
                "地图滤镜",
                "一张供应链地图切换安全、产业、外交镜头。",
                "一次只显示一种镜头，避免信息叠加。",
            ],
            [
                "政策卡",
                "每轮 3 张主卡，其他行动进入抽屉；卡牌带成本和公开信号。",
                "随机性服务情景变化，不制造无意义抽卡。",
            ],
            [
                "科技树",
                "设计、制造、设备、封装、HBM 五条能力路径，研发延迟兑现。",
                "限制同时进行的项目数量，避免资源管理压过博弈。",
            ],
            [
                "城邦外交",
                "荷兰、日本、韩国通过市场准入、材料承诺和安全合作争取。",
                "不需要单位移动，只保留外交筹码和政策交换。",
            ],
            [
                "时代目标",
                "每两轮设置一个阶段任务，完成获得额外政策槽或外交筹码。",
                "目标数量控制在 1 至 2 个。",
            ],
            [
                "危机事件",
                "先进封装事故、出口漏洞、盟友选举、HBM 短缺。",
                "每局只选少量事件，避免随机噪声掩盖战略。",
            ],
            [
                "多结局",
                "技术自主、产业韧性、联盟主导、稳定共存。",
                "首版先做 3 个，确保玩家能提前看见路径。",
            ],
        ],
        [1.25, 3.3, 2.25],
        font_size=8.7,
    )
    doc.add_paragraph(
        "借鉴这些系统的目的，是让玩家的选择形成长期路径，而不是把芯片博弈改造成完整 4X 游戏。"
        "城市微操、单位移动、超大科技树和连续生产队列都会稀释政策判断，首版不应加入。"
    )
    doc.add_heading("首版取舍", level=2)
    add_table(
        doc,
        ["选择", "机制", "理由"],
        [
            [
                "优先加入",
                "地图滤镜、政策卡、第三方外交",
                "直接改善信息读取和战略取舍，与当前模型兼容。",
            ],
            [
                "后续加入",
                "科技树、时代目标、多结局",
                "能增加长期路径，但需要更多内容和平衡。",
            ],
            [
                "首版不加入",
                "城市微操、单位移动、超大科技树",
                "管理成本会压过博弈判断，偏离核心体验。",
            ],
        ],
        [1.25, 2.75, 2.8],
        font_size=8.8,
    )

    add_page_break(doc)

    # Page 8
    doc.add_heading("七  实施路线", level=1)
    doc.add_paragraph(
        "改版应按信息可读性、博弈推理、游戏系统三层推进。每一阶段都以玩家能否完成特定任务"
        "作为验收标准，而不是以界面是否新增模块作为完成标准。"
    )
    add_caption(doc, "表 10  分阶段改版路线", before=True)
    add_table(
        doc,
        ["阶段", "目标", "重点交付", "验收标准"],
        [
            [
                "P0  界面可读性",
                "让玩家在一轮内看懂局势和取舍。",
                "趋势与目标、角色指标、前三影响战报、3 张主卡、隐藏开发标签。",
                "新玩家 60 秒内能说出本局目标、当前瓶颈和一个政策取舍。",
            ],
            [
                "P1  博弈推理",
                "让不完全信息成为可操作的玩法。",
                "稳定隐藏类型、公开信号、信念面板、预测题、可信度、分支预测、复盘。",
                "玩家能说出“我判断对手是哪一型、依据是什么、为何选择这套政策”。",
            ],
            [
                "P2  棋盘化扩展",
                "提高重玩价值和长期策略空间。",
                "地图镜头、第三方外交、简化科技树、阶段目标、危机事件和多结局。",
                "不同开局能形成明显不同的政策路径，且玩家愿意继续下一局。",
            ],
        ],
        [1.28, 2.05, 2.35, 2.15],
        font_size=8.8,
    )

    doc.add_heading("建议的验收指标", level=2)
    reset_numbering()
    add_numbered(doc, "主屏同时出现的战略指标不超过 3 个，主政策卡不超过 3 张。")
    add_numbered(doc, "每轮战报的首屏能明确指出三个变化及其对下一轮的影响。")
    add_numbered(doc, "每轮只有一条核心预测题，避免把推理变成额外表单。")
    add_numbered(doc, "所有信号都必须改变至少一个概率、一个行动选项或一个对手反应。")
    add_numbered(doc, "在 390 像素宽度的移动端上，行动栏不遮挡关键数据，文字不重叠。")

    doc.add_heading("主要风险", level=2)
    add_lead_paragraph(
        doc,
        "对手行为随机化。",
        "如果 AI 类型与行为规则不一致，玩家会把不确定性理解为不公平，而不是风险判断。",
    )
    add_lead_paragraph(
        doc,
        "信号只是装饰。",
        "如果信号不影响后验、可信度或行动，玩家很快会忽略它。系统必须明确哪些判断被改变。",
    )
    add_lead_paragraph(
        doc,
        "系统一次增加过多。",
        "地图、科技树、外交和危机不能同时压到主屏。每轮保持 3 项指标、3 张卡和 1 条预测题。",
    )
    add_lead_paragraph(
        doc,
        "数据与叙事不一致。",
        "趋势、目标、战报和企业利润必须来自同一结算结果，避免界面解释与规则引擎脱节。",
    )

    add_page_break(doc)
    doc.add_heading("八  决策建议", level=1)
    doc.add_paragraph(
        "先做 P0，把数据变化、因果摘要和政策取舍讲清楚，再接入 P1 的信号与信念系统。"
        "P0 和 P1 足以验证核心玩法是否成立；确认玩家愿意观察、判断和复盘后，再投入"
        "地图镜头、外交、科技树和多结局等内容。"
    )
    doc.add_heading("最小原型范围", level=2)
    add_bullet(doc, "只保留中国政府一个可玩角色，美国政府使用两种稳定隐藏类型。")
    add_bullet(doc, "每轮只给 3 张主政策卡，其中至少一张能释放公开信号。")
    add_bullet(doc, "每轮只有一条预测题，结算后展示信念更新和预测误差。")
    add_bullet(doc, "战报首屏只保留三个关键变化，完整记录放入抽屉。")
    add_bullet(doc, "使用现有供应链结算与因果记录，不先扩展地图和科技树。")

    doc.add_heading("首轮可用性测试", level=2)
    add_caption(doc, "表 11  首轮可用性测试任务与通过标准", before=True)
    add_table(
        doc,
        ["测试任务", "观察重点", "通过标准"],
        [
            [
                "说明本局目标与当前瓶颈",
                "玩家是否能在不展开详情的情况下找到方向。",
                "60 秒内说出目标、瓶颈和一项主要代价。",
            ],
            [
                "解释三项指标的环比变化",
                "趋势、目标和变化原因是否可读。",
                "至少正确解释两项变化及其可能来源。",
            ],
            [
                "完成一轮对手判断与政策选择",
                "信号、概率、政策卡和分支预测是否形成一条推理链。",
                "能说明判断概率、依据和选择理由。",
            ],
            [
                "阅读战报并制定下一步",
                "前三影响是否足以支持下一轮行动。",
                "能指出一个关键后果和一个下一轮动作。",
            ],
            [
                "在移动端完成同一轮",
                "固定底栏、横向卡片和文字换行是否正常。",
                "无遮挡、无重叠、无需反复缩放。",
            ],
        ],
        [2.25, 2.45, 2.1],
        font_size=8.8,
    )
    doc.add_paragraph(
        "若玩家能稳定完成上述任务，再进入地图、外交和科技树扩展。"
        "若不能，继续修改信息层级和战报，而不是依靠新增内容弥补理解问题。"
    )

    doc.save(DOCX_PATH)


if __name__ == "__main__":
    ensure_dirs()
    build_docx()
    print(DOCX_PATH)
