from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "output"
DOCX_PATH = OUT_DIR / "芯片外交V1.3修改设计方案.docx"

FONT = "Microsoft YaHei"
BLACK = RGBColor(0, 0, 0)
INK = RGBColor(34, 41, 45)
MUTED = RGBColor(91, 103, 108)
WHITE = RGBColor(255, 255, 255)
HEADER_FILL = "203B46"
ALT_FILL = "F2F5F6"
BORDER_FILL = "D9D9D9"


def set_run_font(run, size: float | None = None, bold: bool | None = None, color=None) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), BORDER_FILL)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_table_layout(table, widths_inches: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(sum(widths_inches) * 1440)))
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    grid = table._tbl.tblGrid
    for index, width in enumerate(widths_inches):
        grid.gridCol_lst[index].set(qn("w:w"), str(int(width * 1440)))
        for cell in table.columns[index].cells:
            cell.width = Inches(width)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    font_size: float = 8.8,
    center_columns: set[int] | None = None,
    header_fill: str = HEADER_FILL,
):
    center_columns = center_columns or set()
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    set_table_layout(table, widths)

    header = table.rows[0]
    set_repeat_table_header(header)
    set_row_cant_split(header)
    for index, text in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        run = paragraph.add_run(text)
        set_run_font(run, font_size, True, WHITE)

    for row_index, values in enumerate(rows):
        row = table.add_row()
        set_row_cant_split(row)
        for col_index, text in enumerate(values):
            cell = row.cells[col_index]
            if row_index % 2 == 1:
                set_cell_shading(cell, ALT_FILL)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if col_index in center_columns else WD_ALIGN_PARAGRAPH.LEFT
            )
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.12
            run = paragraph.add_run(text)
            set_run_font(run, font_size, False, INK)

    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run()
    set_run_font(run, 4)
    return table


def add_body(doc: Document, text: str, space_after: float = 6) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.28
    paragraph.paragraph_format.widow_control = True
    run = paragraph.add_run(text)
    set_run_font(run, 10.5, False, INK)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Inches(0.24)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        paragraph.paragraph_format.space_after = Pt(2.5)
        paragraph.paragraph_format.line_spacing = 1.18
        run = paragraph.add_run(item)
        set_run_font(run, 10.2, False, INK)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.left_indent = Inches(0.26)
        paragraph.paragraph_format.first_line_indent = Inches(-0.2)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.25
        run = paragraph.add_run(item)
        set_run_font(run, 10.2, False, INK)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    if level == 1:
        paragraph.paragraph_format.space_before = Pt(16)
        paragraph.paragraph_format.space_after = Pt(8)
    elif level == 2:
        paragraph.paragraph_format.space_before = Pt(11)
        paragraph.paragraph_format.space_after = Pt(4)
    for run in paragraph.runs:
        set_run_font(run, 15.5 if level == 1 else 11.8, True, BLACK)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.32

    for style_name, size in (("Title", 25), ("Subtitle", 12), ("Heading 1", 15.5), ("Heading 2", 11.8)):
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = BLACK

    title_style = doc.styles["Title"]
    title_p_pr = title_style._element.get_or_add_pPr()
    for border in title_p_pr.findall(qn("w:pBdr")):
        title_p_pr.remove(border)


def add_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(54)

    eyebrow = doc.add_paragraph()
    eyebrow.paragraph_format.space_after = Pt(10)
    run = eyebrow.add_run("产品规则与实现基线")
    set_run_font(run, 10, True, MUTED)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.line_spacing = 1.12
    run = title.add_run("芯片外交 V1.3 修改设计方案")
    set_run_font(run, 26, True, BLACK)

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.paragraph_format.space_after = Pt(34)
    run = subtitle.add_run("删除 ASML S3 并接入灰色交付机制")
    set_run_font(run, 12.5, False, INK)

    meta_rows = [
        ["方案日期", "2026-10-30"],
        ["对照版本", "线上 V1.0 与规则稿 V1.2"],
        ["目标版本", "V1.3"],
        ["文档用途", "产品规则、前端、规则引擎、AI 提示词与测试的共同基线"],
    ]
    add_table(
        doc,
        ["项目", "内容"],
        meta_rows,
        [1.25, 5.85],
        font_size=9.4,
        center_columns={0},
    )

    add_heading(doc, "核心结论", level=2)
    add_body(
        doc,
        "按现有动作编码，待删除项是 ASML 的 S3“维护现有合规客户”，荷兰政府动作 N1、N2、N3 保持不变。"
        "线上 V1.0 的 ASML 只有 S1、S2、S3，而规则稿 V1.2 已设计 S4 灰色渠道；只删 S3 会让 ASML 只剩两项行动，"
        "压缩谈判与反制空间，因此建议把“停用 S3”和“启用 S4”放在同一次版本发布中。",
    )
    add_body(
        doc,
        "V1.3 继续采用固定四回合、秘密行动、信用和协议履约机制。变更重点不是增加一个高风险按钮，而是让 ASML 在"
        "合规订单、市场转移和灰色交付之间形成可识别、可承担的取舍。",
    )
    add_bullets(
        doc,
        [
            "ASML 合法动作调整为 S1、S2、S4 三项，另可选择 PASS。",
            "美国政府增加 U4 跨境合规执法，灰色交付按概率查获，未查获时不公开具体动作。",
            "保留线上 A、B、H、R 四个指标字段，规则稿的 T、A、E、S 只用于解释语义，避免同字母异义。",
            "所有动作效果统一为整数；先合计增量，再一次性截断到 0 至 5。",
        ],
    )

    doc.add_page_break()


def build_document() -> Document:
    doc = Document()
    configure_styles(doc)

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    add_cover(doc)

    add_heading(doc, "一 版本基线与总体决策", level=1)
    add_body(
        doc,
        "线上 V1.0 已具备四人角色、固定四回合、私人目标、正式协议、信用和履约记录。规则稿 V1.2 增加了灰色交付、"
        "美国跨境执法和两级胜利，但动作表与线上版本并不一致。V1.3 的目标是先完成动作集和灰色渠道的统一，再讨论"
        "更复杂的长期模型。",
    )
    add_table(
        doc,
        ["模块", "线上 V1.0", "规则稿 V1.2", "V1.3 建议"],
        [
            ["回合结构", "固定 4 回合", "固定 4 回合", "保持固定 4 回合"],
            ["ASML 动作", "S1、S2、S3", "S1、S2、S3、S4", "S1、S2、S4；删除 S3"],
            ["美国动作", "U1、U2、U3", "U1、U2、U3、U4", "增加 U4"],
            ["灰色交付", "未实现", "有查获概率和处罚", "按 V1.2 机制实现"],
            ["指标表达", "A/B/H/R 整数", "T/A/E/S 含小数效果", "保留 A/B/H/R，效果改为整数"],
            ["胜负表达", "正式胜利", "实质胜利与外交胜利", "先实现两级结果标签"],
            ["公开行动", "结算后公开实际行动", "未查获 S4 不公开", "改为只公开可观察行动"],
        ],
        [1.0, 1.65, 1.85, 2.2],
        font_size=8.4,
        center_columns={0},
    )

    add_heading(doc, "本次发布的三项硬约束", level=2)
    add_numbered(
        doc,
        [
            "S3 必须从规则引擎的合法动作白名单中删除，不能只隐藏页面按钮。",
            "S4 与 U4 必须在同一版本启用，否则 ASML 的行动池会过窄。",
            "进行中的 V1.0 对局保持原版本只读；新对局使用 V1.3 动作集，不做 S3 到 S4 的静默映射。",
        ],
    )

    add_heading(doc, "二 动作表修改", level=1)
    add_body(
        doc,
        "下表是 V1.3 建议动作集。效果继续使用线上 A、B、H、R 字段：A 为技术合作机会，B 为商业活力，"
        "H 为美荷协调，R 为中国自主技术。规则稿中的“技术管制强度”与 A 的变化方向相反，不能直接同名迁移。",
    )
    add_table(
        doc,
        ["角色", "编号", "行动", "合法条件", "建议效果", "状态"],
        [
            ["美国", "U1", "收紧技术管制", "无", "A-1，B-1", "保留"],
            ["美国", "U2", "有限放宽管制", "无", "A+1，B+1", "保留"],
            ["美国", "U3", "与荷兰开展产业合作", "无", "B+1，H+1", "保留"],
            ["美国", "U4", "加强跨境合规执法", "无", "B-1；本轮 S4 查获率+30%；查获后 H+1", "新增"],
            ["中国", "C1", "投资自主研发", "最多 2 次；第 4 回合不可选", "下回合开始 R+1；本回合 B-1", "调整"],
            ["中国", "C2", "限制关键材料出口", "无", "B-2", "保留"],
            ["中国", "C3", "提供市场合作便利", "无", "B+1", "保留"],
            ["荷兰", "N1", "收紧设备许可", "无", "A-1，B-1，H+1", "保留"],
            ["荷兰", "N2", "批准有限许可", "回合开始时 H≤3", "H-1；若 ASML 同轮选 S1，再触发 A+1、B+2", "保留"],
            ["荷兰", "N3", "提议美荷联合项目", "无", "无基础效果；若美国同轮选 U3，再触发 B+1、H+1", "保留"],
            ["ASML", "S1", "执行对华合规订单", "无", "荷兰同轮选 N2：A+1、B+2；否则 B-1", "保留"],
            ["ASML", "S2", "拓展其他市场", "无", "B+1，A-1", "保留"],
            ["ASML", "S3", "维护现有合规客户", "无", "原效果停用", "删除"],
            ["ASML", "S4", "灰色渠道交付", "无需 N2 前置许可", "未查获 A+1、B+3；查获 B-2、credit-1、下回合 A-1", "新增"],
            ["全体", "PASS", "维持现状", "无", "无指标变化；协议不得承诺 PASS", "保留"],
        ],
        [0.58, 0.56, 1.42, 1.55, 2.42, 0.58],
        font_size=7.8,
        center_columns={0, 1, 5},
    )
    add_body(
        doc,
        "H 对政府技术合作变化仍起杠杆作用：U1、U2、N1 及 N2 与 S1 联动产生的 A 变化，按变化幅度乘以"
        "一加五分之一 H，再对幅度向下取整并保留正负号。ASML 自身 S2 与 S4 产生的 A 变化不受该杠杆放大。",
    )

    add_heading(doc, "三 S3 删除的实施范围", level=1)
    add_body(
        doc,
        "S3 的设计效果取决于回合开始时的 A 值，与 PASS 和 S1 的边界接近，缺少稳定的谈判价值。删除后不需要把原效果"
        "迁移到 S1 或其他动作；V1.3 用 S4 提供第三种明确策略。任何页面文案、示例或测试夹具中都不得继续把 S3 列为"
        "可选动作。",
    )
    add_table(
        doc,
        ["修改面", "必须完成的内容", "失败表现"],
        [
            ["动作定义", "ASML 动作数组只保留 S1、S2、S4；名称 W(S3) 不再由新对局解析。", "页面仍显示 S3，或 API 返回 S3。"],
            ["服务端校验", "行动提交、协议承诺和 AI 输出校验都拒绝 S3，并返回稳定错误码。", "绕过前端后仍能提交 S3。"],
            ["协议系统", "新协议不能承诺 S3；历史协议可保留 S3 文本但不能恢复执行。", "签约时报错或静默改成其他动作。"],
            ["AI Agent", "系统提示词、工具描述和合法动作列表均不含 S3；非法输出重试一次后回退 PASS。", "Agent 继续选择 S3。"],
            ["前端界面", "ASML 行动区显示三张卡；不再出现“维护现有合规客户”。", "空按钮、重复卡或布局跳动。"],
            ["战报", "新战报不出现 S3；旧对局可显示“已停用动作”并保持只读。", "回放内容被篡改或新旧语义混用。"],
            ["分析埋点", "按 schemaVersion 区分 S3 历史数据与 V1.3 新数据。", "把旧 S3 频率误认为新版本行为。"],
            ["测试与文档", "删除所有 S3 合法路径用例，新增 S3 拒绝用例和旧存档兼容用例。", "测试仍要求 S3 可用。"],
        ],
        [0.95, 4.25, 1.9],
        font_size=8.5,
        center_columns={0},
    )

    add_heading(doc, "四 灰色交付与结算顺序", level=1)
    add_body(
        doc,
        "灰色交付查获率以回合开始时 H 计算：基础 20%，每点 H 增加 10%，美国本轮选择 U4 再增加 30%；"
        "若 ASML 曾因灰色交付被查获并使信用降至 0，后续基础查获率再增加 20%。所有加成相加后截断到 100%。"
        "H=2 且无 U4 时为 40%，有 U4 时为 70%；H=5 且有 U4 时固定为 100%。",
    )
    add_table(
        doc,
        ["阶段", "处理规则", "需要记录的内部数据"],
        [
            ["回合开始", "先执行上回合遗留效果，再保存 roundStart 快照。", "待生效研发、待生效 A 变化、回合初始 H。"],
            ["行动提交", "四方同时提交，提交后不可更改。", "行动 ID、提交时间、schemaVersion。"],
            ["基础增量", "计算 U1 至 U4、C1 至 C3、N1 至 N3、S1、S2、S4 的常规效果。", "动作来源与原始增量。"],
            ["S4 判定", "用 roundStart H 和本轮 U4 计算查获率，进行一次伯努利判定。", "随机种子、查获概率、判定结果，仅内部可见。"],
            ["联动与取整", "应用政府 A 变化杠杆，按符号保留幅度取整，再合计所有增量。", "每项修正后的增量。"],
            ["状态更新", "一次性截断到 0 至 5，更新信用和待生效效果。", "最终状态、信用变化原因。"],
            ["公开战报", "公开可观察行动和指标结果；未查获 S4 不得显示动作名称或灰色渠道解释。", "公开视图与内部审计视图分离。"],
        ],
        [0.9, 4.1, 2.1],
        font_size=8.4,
        center_columns={0},
    )
    add_body(
        doc,
        "规则引擎不得直接计算期望收益后替代随机结果。查获率只用于显示风险和内部验证；每一局必须实际判定一次，"
        "并保存可复现的随机种子。未查获时，公开战报可以显示 A、B 的最终变化，但不能把变化原因写成“灰色交付”。",
    )

    add_heading(doc, "五 规则一致性修订", level=1)
    add_body(
        doc,
        "规则稿中存在几处会直接影响实现的表述冲突。V1.3 应先把这些冲突收敛为唯一口径，再进入前后端联调。",
    )
    add_table(
        doc,
        ["问题", "现有冲突", "V1.3 口径"],
        [
            ["公开行动边界", "一处写公开实际行动，另一处写未查获 S4 不公开。", "统一为“公开可观察行动”；未查获 S4 不进入公共行动日志。"],
            ["整数与小数", "指标规定为整数，但存在 A+0.5、E-0.3 等效果。", "动作效果改为整数；内部不使用小数状态。"],
            ["负向取整", "直接向下取整会把 -1 放大成 -1.4 再变成 -2。", "按幅度向下取整，再恢复正负号。"],
            ["字母重名", "V1.2 的 A 是同盟协调度，线上 A 是技术合作机会。", "保留线上 A/B/H/R；文档同时给出 T/A/E/S 语义映射。"],
            ["协议合法性", "协议可以承诺动作，但动作可能因条件变为非法。", "发起和接受时校验；结算前仍非法则标记冲突，不静默替换。"],
            ["旧对局兼容", "旧状态含 S3，新引擎没有 S3。", "按版本冻结旧对局；新引擎只处理 V1.3，不做动作映射。"],
            ["AI 非法输出", "LLM 可能输出已删除或不存在的动作。", "重试一次；仍非法则回退 PASS，并记录内部事件。"],
            ["随机结果复盘", "只保存最终指标无法还原查获判定。", "保存随机种子、概率和结果，公开层不暴露种子。"],
        ],
        [1.0, 2.75, 3.35],
        font_size=8.4,
        center_columns={0},
    )

    add_heading(doc, "指标语义对照", level=2)
    add_table(
        doc,
        ["规则稿符号", "规则稿含义", "线上字段", "线上含义", "迁移处理"],
        [
            ["T", "技术管制强度", "A", "技术合作机会", "方向相反，按语义换算后写入 A"],
            ["E", "对华产业收益", "B", "商业活力", "保留 B，作为整体商业环境"],
            ["A", "同盟协调度", "H", "美荷协调", "保留 H，避免与线上 A 冲突"],
            ["S", "中国自主能力", "R", "中国自主技术", "保留 R，字段语义直接对应"],
        ],
        [0.85, 1.45, 0.75, 1.45, 2.6],
        font_size=8.6,
        center_columns={0, 2},
    )

    doc.add_page_break()
    add_heading(doc, "六 目标与胜利表达", level=1)
    add_body(
        doc,
        "当前线上采用单一的“正式胜利”，要求目标阈值、至少一份全员履约协议和信用不低于 1。V1.3 可以保留这一完整胜利，"
        "增加“实质胜利”作为较低层级的结果标签：数值达标且信用不低于 1，但不要求履约协议。这样既能保留谈判价值，也能让"
        "单边达成目标的玩家看到结果。",
    )
    add_table(
        doc,
        ["角色", "类型", "线上 A/B/H/R 条件", "结果标签"],
        [
            ["美国", "安全优先", "A≤1 且 R≤2", "实质胜利"],
            ["美国", "联盟优先", "H≥4 且 A≤2", "实质胜利"],
            ["中国", "自主突破", "R≥3 且 A≥2", "实质胜利"],
            ["中国", "合作维持", "A≥3 且 B≥4", "实质胜利"],
            ["荷兰", "协调优先", "H≥4 且 B≥3", "实质胜利"],
            ["荷兰", "商业优先", "A≥3 且 B≥4", "实质胜利"],
            ["ASML", "对华深耕", "A≥3 且 B≥4", "实质胜利"],
            ["ASML", "全球多元", "B=5 且 H≥3", "实质胜利"],
        ],
        [0.9, 1.35, 3.15, 1.7],
        font_size=8.7,
        center_columns={0, 3},
    )
    add_body(
        doc,
        "最终结算时按顺序检查：条件与信用是否满足，决定是否获得实质胜利；再检查是否至少有一份本人参与且全员履约的协议，"
        "满足时升级为外交胜利。结果页同时显示两种标签、未满足条件、信用和履约协议数量，便于复盘。",
    )

    add_heading(doc, "七 界面与 AI 行为", level=1)
    add_body(
        doc,
        "动作删除后，ASML 行动区应稳定显示 S1、S2、S4 三张卡，不留下空位。S4 卡片必须明确写出查获率区间、未查获和"
        "查获两类结果，以及信用下降后的持续风险。N2 的禁用条件应在签约前显示，而不是等到行动提交时才报错。",
    )
    add_bullets(
        doc,
        [
            "行动卡显示前置条件、基础效果和联动效果，S4 额外显示当前查获率。",
            "协议卡只允许选择当前合法动作；H≥4 时 N2 不可选，并说明原因。",
            "战报对未查获 S4 只展示公开指标变化；对查获 S4 展示处罚、信用变化和下回合 A 变化。",
            "结果页显示实质胜利或外交胜利，不把所有未获外交胜利的玩家统一写成失败。",
            "AI Agent 与秘书 Agent 共用服务端合法动作和可见信息边界，不读取隐藏判定或未查获 S4 记录。",
        ],
    )

    add_heading(doc, "八 实施顺序", level=1)
    add_table(
        doc,
        ["阶段", "交付内容", "完成标准"],
        [
            ["1 冻结规则", "确定 V1.3 动作表、整数效果、指标映射和胜负层级。", "产品规则、界面文案和规则引擎使用同一动作清单。"],
            ["2 升级引擎", "新增 schemaVersion、U4、S4、查获判定、信用处罚和旧对局只读。", "接口无法提交 S3；旧对局不被新引擎改写。"],
            ["3 更新界面", "ASML 三卡布局、S4 风险说明、N2 条件、两级结果和隐藏战报。", "桌面与移动端均无空卡、溢出或信息泄露。"],
            ["4 对齐 AI", "更新 Agent 工具、提示词、协议校验、非法输出回退。", "Agent 无法选择 S3，且不会读取隐藏结果。"],
            ["5 回归测试", "动作、概率、信用、协议、战报、存档和回放测试。", "下表验收项全部通过。"],
            ["6 平衡测试", "规则型 Agent 批量自博弈，再抽样真实 LLM 对局。", "无明显单角色压制或行动频率异常。"],
        ],
        [0.95, 3.0, 3.15],
        font_size=8.6,
        center_columns={0},
    )

    add_heading(doc, "九 验收与平衡测试", level=1)
    add_table(
        doc,
        ["编号", "测试项", "预期结果"],
        [
            ["T01", "ASML 新对局合法动作", "仅返回 S1、S2、S4 和 PASS，不含 S3。"],
            ["T02", "提交 S3", "服务端返回明确非法动作，不改变状态。"],
            ["T03", "协议承诺 S3", "发起或接受阶段被拒绝，不生成生效协议。"],
            ["T04", "AI 输出 S3", "重试一次后回退 PASS，内部记录原因。"],
            ["T05", "旧对局含 S3", "可查看和回放，不能继续行动，也不自动映射为 S4。"],
            ["T06", "S4 未查获", "A+1、B+3，公共日志不出现 S4 名称或灰色渠道解释。"],
            ["T07", "S4 被查获", "B-2、credit-1，下回合开始 A-1，并公开处罚结果。"],
            ["T08", "查获概率", "H=2 无 U4 为 40%，有 U4 为 70%；H=5 有 U4 为 100%。"],
            ["T09", "N2 条件", "回合开始时 H≥4 时不可选；H=3 时可选。"],
            ["T10", "整数结算", "任一动作组合的最终状态均为 0 至 5 的整数。"],
            ["T11", "复盘一致性", "同一种子重复结算得到同一查获结果和指标状态。"],
            ["T12", "结果层级", "实质胜利与外交胜利标签、条件说明和信用记录一致。"],
        ],
        [0.5, 1.8, 4.8],
        font_size=8.7,
        center_columns={0},
    )
    add_body(
        doc,
        "首轮平衡测试至少记录各角色胜率、各动作使用率、S4 选择率、平均查获率、协议形成率、违约率、N2 因 H 条件被禁用的次数"
        "和无胜者比例。先用规则型 Agent 跑大批量机械回归，再用少量真实 LLM 对局检查谈判是否围绕 S4、N2 和履约协议展开。",
    )

    add_heading(doc, "十 可进一步修改的地方", level=1)
    add_body(
        doc,
        "以下项目不建议与 S3 删除同时上线。先在 V1.3 获得稳定的动作频率和胜率数据，再决定是否扩大规则复杂度。",
    )
    add_table(
        doc,
        ["优先级", "修改方向", "建议做法", "判断依据"],
        [
            ["P1", "自研对冲", "R≥3 时减弱技术收紧造成的 B 损失；R≥4 时按规则调整有效技术压制。", "若美国收紧策略胜率明显过高，再启用。"],
            ["P1", "动作预览", "行动提交前显示基础效果、联动条件和最坏结果，不预测对手选择。", "降低新玩家误解 S1 与 S4 的成本。"],
            ["P1", "S4 风险曲线", "根据首轮自博弈调整基础查获率或 U4 加成。", "若 S4 使用率过高或几乎无人使用。"],
            ["P2", "跨回合协议", "允许协议承诺两回合内的行动组合，并明确违约时点。", "在四方谈判已经稳定后进行。"],
            ["P2", "第三方转口", "增加一条成本更高、查获后处罚更重的灰色路径。", "当 S4 已被玩家充分理解后。"],
            ["P2", "随机事件卡", "加入供应链扰动或国内政治事件，限制单次波动幅度。", "避免事件压过玩家行动的主因果。"],
            ["P3", "判断力计分", "在结果页评价对手类型判断和协议判断，独立于胜负。", "用于教学和复盘，不改变本局胜负。"],
            ["P3", "扩展角色", "引入日本或台积电作为环境变量，而非直接增加第五名玩家。", "先验证四角色版本的稳定性和时长。"],
        ],
        [0.6, 1.25, 3.3, 1.95],
        font_size=8.35,
        center_columns={0},
    )

    doc.add_page_break()
    add_heading(doc, "十一 上线检查表", level=1)
    add_table(
        doc,
        ["检查项", "结果"],
        [
            ["新对局动作表中不存在 S3", "待勾选"],
            ["U4 与 S4 在同一 schemaVersion 启用", "待勾选"],
            ["协议、AI、API 均拒绝 S3", "待勾选"],
            ["旧对局只读且可回放", "待勾选"],
            ["未查获 S4 不出现在公共行动日志", "待勾选"],
            ["查获率、信用、下回合效果和随机种子可复现", "待勾选"],
            ["N2 在 H≥4 时的禁用文案已上线", "待勾选"],
            ["结果页支持实质胜利与外交胜利", "待勾选"],
            ["桌面端和移动端 ASML 三卡布局通过截图检查", "待勾选"],
            ["规则型自博弈与真实 LLM 对局完成抽样复盘", "待勾选"],
        ],
        [5.35, 1.4],
        font_size=9.0,
        center_columns={1},
    )
    add_body(
        doc,
        "发布后先观察动作选择、协议履约和异常报错，再决定是否进入自研对冲、跨回合协议或新增灰色路径。"
        "如果在首轮数据中 S4 被高频无脑选择，应优先调整查获率或 U4 成本，而不是继续增加新动作。",
    )

    return doc


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.core_properties.title = "芯片外交 V1.3 修改设计方案"
    doc.core_properties.subject = "删除 ASML S3 并接入灰色交付机制"
    doc.core_properties.author = "芯片外交项目组"
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
