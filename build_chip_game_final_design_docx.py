from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

import build_game_improvement_docx as base


ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "output"
ASSET_DIR = OUT_DIR / "_assets_final_design"
QA_DIR = OUT_DIR / "_qa_final_design"
DOCX_PATH = OUT_DIR / "芯片战争_不完全信息动态博弈_最终设计方案.docx"

BODY_FONT = base.BODY_FONT
INK = base.INK
MUTED = base.MUTED
BLACK = base.BLACK
WHITE = base.WHITE


def section_heading(doc: Document, text: str, level: int = 1, page_break: bool = True):
    heading = doc.add_heading(text, level=level)
    if page_break:
        heading.paragraph_format.page_break_before = True
    return heading


def add_footer(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("芯片战争不完全信息动态博弈最终游戏设计方案  |  第 ")
    run.font.name = BODY_FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED

    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    field_separate = OxmlElement("w:fldChar")
    field_separate.set(qn("w:fldCharType"), "separate")
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instr_text)
    run._r.append(field_separate)
    run._r.append(field_end)

    suffix = paragraph.add_run(" 页")
    suffix.font.name = BODY_FONT
    suffix._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
    suffix.font.size = Pt(8)
    suffix.font.color.rgb = MUTED


def build_timeline(path: Path) -> None:
    scale = 2
    width, height = 1600 * scale, 720 * scale
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)

    title_font = base.pil_font(38 * scale, bold=True)
    body_font = base.pil_font(21 * scale)
    num_font = base.pil_font(32 * scale, bold=True)
    small_font = base.pil_font(18 * scale)

    draw.text((65 * scale, 35 * scale), "六轮结构与选举冲击", font=title_font, fill="#17262C")
    draw.text(
        (65 * scale, 93 * scale),
        "前三轮判断旧政府，第三轮结束后进入新政府过渡期",
        font=body_font,
        fill="#5B6D74",
    )

    box_w = 180 * scale
    box_h = 120 * scale
    gap = 26 * scale
    start_x = 65 * scale
    top_y = 210 * scale

    stages = [
        ("R1", "旧政府", "类型稳定"),
        ("R2", "旧政府", "信号累积"),
        ("R3", "旧政府", "大选前夜"),
        ("换台", "执政联盟更替", "新类型未知"),
        ("R4", "过渡政府", "政策惯性"),
        ("R5", "新政府", "行为定型"),
        ("R6", "新政府", "最终结算"),
    ]
    centers: list[tuple[int, int]] = []

    for index, (round_text, stage, subtext) in enumerate(stages):
        x1 = start_x + index * (box_w + gap)
        x2 = x1 + box_w
        y1 = top_y
        y2 = y1 + box_h
        if index <= 2:
            fill, outline = "#EBF2F4", "#9EB2BA"
        elif index == 3:
            fill, outline = "#F6E9DA", "#C58B38"
        elif index == 4:
            fill, outline = "#E6ECEE", "#7F949C"
        else:
            fill, outline = "#EDF4F2", "#7F9D94"
        draw.rounded_rectangle(
            (x1, y1, x2, y2),
            radius=18 * scale,
            fill=fill,
            outline=outline,
            width=3 * scale,
        )
        draw_round = round_text if round_text != "换台" else "!"
        base.draw_centered_text(
            draw,
            (x1, y1 + 10 * scale, x2, y1 + 54 * scale),
            draw_round,
            num_font,
            "#17262C",
        )
        base.draw_centered_text(
            draw,
            (x1 + 8 * scale, y1 + 58 * scale, x2 - 8 * scale, y1 + 83 * scale),
            stage,
            body_font,
            "#17262C",
        )
        base.draw_centered_text(
            draw,
            (x1 + 8 * scale, y1 + 87 * scale, x2 - 8 * scale, y1 + 111 * scale),
            subtext,
            small_font,
            "#65777E",
        )
        centers.append(((x1 + x2) // 2, (y1 + y2) // 2))
        if index < len(stages) - 1:
            base.draw_arrow(
                draw,
                (x2 + 2 * scale, y1 + box_h // 2),
                (x2 + gap - 3 * scale, y1 + box_h // 2),
                fill="#8EA1A9",
                width=5 * scale,
            )

    band_y = 415 * scale
    bands = [
        (82 * scale, 836 * scale, "#E9F1F4", "旧类型：持续观察与验证"),
        (878 * scale, 1225 * scale, "#F7EEE2", "过渡：政策惯性与新证据"),
        (1267 * scale, 1518 * scale, "#EAF2EF", "新类型：重新判断"),
    ]
    for x1, x2, fill, label in bands:
        draw.rounded_rectangle(
            (x1, band_y, x2, band_y + 82 * scale),
            radius=14 * scale,
            fill=fill,
            outline="#D3DDE0",
            width=2 * scale,
        )
        base.draw_centered_text(
            draw,
            (x1 + 12 * scale, band_y, x2 - 12 * scale, band_y + 82 * scale),
            label,
            body_font,
            "#4F6269",
        )

    draw.text(
        (82 * scale, 545 * scale),
        "旧证据仍保留，但大选后的类型先验会被重置并重新更新。",
        font=body_font,
        fill="#5B6D74",
    )
    image.save(path, quality=95)


def build_ui(path: Path) -> None:
    scale = 2
    width, height = 1600 * scale, 1120 * scale
    image = Image.new("RGB", (width, height), "#F3F6F7")
    draw = ImageDraw.Draw(image)

    title_font = base.pil_font(34 * scale, bold=True)
    sub_font = base.pil_font(22 * scale, bold=True)
    body_font = base.pil_font(19 * scale)
    small_font = base.pil_font(16 * scale)
    metric_font = base.pil_font(28 * scale, bold=True)

    draw.rounded_rectangle(
        (35 * scale, 28 * scale, 1565 * scale, 1090 * scale),
        radius=22 * scale,
        fill="#FFFFFF",
        outline="#B8C5CA",
        width=3 * scale,
    )
    draw.rounded_rectangle(
        (55 * scale, 50 * scale, 1545 * scale, 175 * scale),
        radius=14 * scale,
        fill="#E7EFF1",
        outline="#C6D1D5",
        width=2 * scale,
    )
    draw.text((78 * scale, 73 * scale), "第04轮 / 新政府过渡期", font=title_font, fill="#17262C")
    draw.text(
        (78 * scale, 122 * scale),
        "阶段任务：稳定国产交付，同时避免产业压力失控",
        font=body_font,
        fill="#596B72",
    )

    metrics = [
        ("可用算力", "49", "+6.5"),
        ("自主度", "62", "B"),
        ("产业压力", "46", "+12"),
    ]
    for index, (label, value, delta) in enumerate(metrics):
        x1 = 840 * scale + index * 225 * scale
        x2 = x1 + 205 * scale
        draw.rounded_rectangle(
            (x1, 65 * scale, x2, 155 * scale),
            radius=12 * scale,
            fill="#FFFFFF",
            outline="#CAD4D7",
            width=2 * scale,
        )
        draw.text((x1 + 14 * scale, 78 * scale), label, font=small_font, fill="#66787F")
        draw.text((x1 + 14 * scale, 103 * scale), value, font=metric_font, fill="#17262C")
        draw.text((x1 + 128 * scale, 109 * scale), delta, font=small_font, fill="#9B671C")

    draw.rounded_rectangle(
        (55 * scale, 195 * scale, 1095 * scale, 805 * scale),
        radius=16 * scale,
        fill="#FFFFFF",
        outline="#C5D0D4",
        width=2 * scale,
    )
    draw.rounded_rectangle(
        (1115 * scale, 195 * scale, 1545 * scale, 805 * scale),
        radius=16 * scale,
        fill="#FFFFFF",
        outline="#C5D0D4",
        width=2 * scale,
    )
    draw.text((78 * scale, 218 * scale), "供应链地图", font=title_font, fill="#17262C")
    draw.text((78 * scale, 262 * scale), "当前镜头：产业", font=small_font, fill="#66787F")

    nodes = [
        ("芯片设计", "38.9", "稳定"),
        ("工艺良率", "32.2", "承压"),
        ("制造设备", "36.0", "稳定"),
        ("先进封装", "25.0", "瓶颈"),
        ("HBM 配套", "29.0", "偏紧"),
    ]
    node_w = 178 * scale
    node_h = 178 * scale
    gap = 30 * scale
    y1 = 390 * scale
    for index, (name, value, status) in enumerate(nodes):
        x1 = 82 * scale + index * (node_w + gap)
        x2 = x1 + node_w
        y2 = y1 + node_h
        fill = "#F6EADD" if index == 3 else "#EDF3F5"
        outline = "#C18B3B" if index == 3 else "#A7B8BF"
        draw.rounded_rectangle(
            (x1, y1, x2, y2),
            radius=15 * scale,
            fill=fill,
            outline=outline,
            width=3 * scale,
        )
        base.draw_centered_text(
            draw,
            (x1 + 8 * scale, y1 + 18 * scale, x2 - 8 * scale, y1 + 46 * scale),
            name,
            sub_font,
            "#17262C",
        )
        base.draw_centered_text(
            draw,
            (x1 + 8 * scale, y1 + 72 * scale, x2 - 8 * scale, y1 + 120 * scale),
            value,
            metric_font,
            "#17262C",
        )
        base.draw_centered_text(
            draw,
            (x1 + 8 * scale, y1 + 136 * scale, x2 - 8 * scale, y1 + 166 * scale),
            status,
            small_font,
            "#8C5F1B" if index == 3 else "#60737A",
        )
        if index < len(nodes) - 1:
            base.draw_arrow(
                draw,
                (x2 + 3 * scale, y1 + node_h // 2),
                (x2 + gap - 4 * scale, y1 + node_h // 2),
                fill="#8EA1A9",
                width=5 * scale,
            )

    draw.text((82 * scale, 632 * scale), "当前瓶颈：先进封装", font=sub_font, fill="#8C5F1B")
    draw.text(
        (82 * scale, 672 * scale),
        "点击节点查看库存、替代供给、企业暴露与三回合趋势。",
        font=body_font,
        fill="#566970",
    )
    draw.text(
        (82 * scale, 712 * scale),
        "大选换台后，旧证据仍保留；新政府证据优先影响当前判断。",
        font=body_font,
        fill="#566970",
    )

    draw.text((1140 * scale, 218 * scale), "情报与对手判断", font=title_font, fill="#17262C")
    draw.text((1140 * scale, 272 * scale), "新政府公开信号", font=sub_font, fill="#17262C")
    draw.text(
        (1140 * scale, 310 * scale),
        "新政府宣布重新审查许可\n来源：政策简报  可信度 70%",
        font=body_font,
        fill="#566970",
    )
    draw.text((1140 * scale, 402 * scale), "当前类型判断", font=sub_font, fill="#17262C")
    probs = [("产业敏感型", "45%", 0.45, "#B87920"), ("条件交换型", "35%", 0.35, "#6E8B95"), ("安全优先型", "20%", 0.20, "#A9BAC0")]
    for index, (label, pct, ratio, color) in enumerate(probs):
        y = 444 * scale + index * 57 * scale
        draw.text((1140 * scale, y), f"{label}  {pct}", font=body_font, fill="#566970")
        draw.rounded_rectangle(
            (1140 * scale, y + 30 * scale, 1511 * scale, y + 46 * scale),
            radius=8 * scale,
            fill="#E7EDEF",
        )
        draw.rounded_rectangle(
            (1140 * scale, y + 30 * scale, 1140 * scale + int(371 * ratio) * scale, y + 46 * scale),
            radius=8 * scale,
            fill=color,
        )

    draw.text((1140 * scale, 622 * scale), "本轮预测", font=sub_font, fill="#17262C")
    draw.text(
        (1140 * scale, 660 * scale),
        "新政府会在第五轮前放松许可吗？\n你的判断：60%",
        font=body_font,
        fill="#566970",
    )
    draw.text(
        (1140 * scale, 730 * scale),
        "判断锁定后，再选择政策。",
        font=small_font,
        fill="#8C5F1B",
    )

    draw.rounded_rectangle(
        (55 * scale, 830 * scale, 1545 * scale, 1045 * scale),
        radius=16 * scale,
        fill="#EAF0F2",
        outline="#C3CFD3",
        width=2 * scale,
    )
    draw.text((78 * scale, 852 * scale), "本轮政策手牌", font=title_font, fill="#17262C")
    draw.text(
        (78 * scale, 898 * scale),
        "2 AP  /  3 张主卡  /  公开表态：模糊",
        font=body_font,
        fill="#596B72",
    )

    cards = [
        ("对日市场准入", "外交", "改善协调\n消耗 15 预算"),
        ("收紧镓许可", "管制", "增强姿态\n产业承压"),
        ("国产封装协作", "产业", "直击瓶颈\n消耗 10 预算"),
    ]
    for index, (name, kind, note) in enumerate(cards):
        x1 = 400 * scale + index * 355 * scale
        x2 = x1 + 325 * scale
        draw.rounded_rectangle(
            (x1, 858 * scale, x2, 1008 * scale),
            radius=13 * scale,
            fill="#FFFFFF",
            outline="#B6C4C9",
            width=2 * scale,
        )
        draw.text((x1 + 16 * scale, 876 * scale), name, font=sub_font, fill="#17262C")
        draw.text((x1 + 16 * scale, 918 * scale), kind, font=small_font, fill="#8C5F1B")
        draw.multiline_text(
            (x1 + 100 * scale, 914 * scale),
            note,
            font=small_font,
            fill="#596B72",
            spacing=6 * scale,
        )

    image.save(path, quality=95)


def build_agent_pipeline(path: Path) -> None:
    scale = 2
    width, height = 1600 * scale, 900 * scale
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)

    title_font = base.pil_font(38 * scale, bold=True)
    body_font = base.pil_font(20 * scale)
    box_font = base.pil_font(22 * scale, bold=True)
    small_font = base.pil_font(16 * scale)

    draw.text((65 * scale, 34 * scale), "Agent 与局部博弈树流程", font=title_font, fill="#17262C")
    draw.text(
        (65 * scale, 92 * scale),
        "规则控制事实，Agent 选择动作，大模型只负责叙述",
        font=body_font,
        fill="#5B6D74",
    )

    boxes = [
        ("观察", "公共状态、证据\n私有类型、记忆"),
        ("信念更新", "类型概率\n证据可信度"),
        ("局部博弈树", "候选动作 × 类型\n× 第三方反应"),
        ("动作选择", "期望效用\n规则校验"),
        ("结算", "政策、企业、供应链\n国家收益"),
        ("后验反馈", "实际行为\n信念与复盘"),
    ]
    box_w = 220 * scale
    box_h = 185 * scale
    gap = 38 * scale
    start_x = 65 * scale
    y1 = 270 * scale
    for index, (name, note) in enumerate(boxes):
        x1 = start_x + index * (box_w + gap)
        x2 = x1 + box_w
        y2 = y1 + box_h
        fill = "#EFF4F5" if index not in (2, 3) else "#F6EBDD"
        outline = "#A4B7BE" if index not in (2, 3) else "#C28A38"
        draw.rounded_rectangle(
            (x1, y1, x2, y2),
            radius=18 * scale,
            fill=fill,
            outline=outline,
            width=3 * scale,
        )
        base.draw_centered_text(
            draw,
            (x1 + 10 * scale, y1 + 22 * scale, x2 - 10 * scale, y1 + 58 * scale),
            name,
            box_font,
            "#17262C",
        )
        base.draw_centered_text(
            draw,
            (x1 + 12 * scale, y1 + 75 * scale, x2 - 12 * scale, y2 - 15 * scale),
            note,
            body_font,
            "#596B72",
        )
        if index < len(boxes) - 1:
            base.draw_arrow(
                draw,
                (x2 + 2 * scale, y1 + box_h // 2),
                (x2 + gap - 3 * scale, y1 + box_h // 2),
                fill="#8EA1A9",
                width=5 * scale,
            )

    llm_box = (760 * scale, 615 * scale, 1230 * scale, 775 * scale)
    draw.rounded_rectangle(
        llm_box,
        radius=16 * scale,
        fill="#F8F3EA",
        outline="#CAA068",
        width=3 * scale,
    )
    base.draw_centered_text(
        draw,
        (780 * scale, 630 * scale, 1210 * scale, 665 * scale),
        "大模型叙事层",
        box_font,
        "#17262C",
    )
    base.draw_centered_text(
        draw,
        (785 * scale, 680 * scale, 1205 * scale, 758 * scale),
        "把结构化事件改写成新闻、讲话或简报\n不决定类型、概率、收益和结算",
        body_font,
        "#596B72",
    )
    draw.line(
        (720 * scale, 695 * scale, 760 * scale, 695 * scale),
        fill="#B98843",
        width=4 * scale,
    )
    draw.line(
        (1230 * scale, 695 * scale, 1270 * scale, 695 * scale),
        fill="#B98843",
        width=4 * scale,
    )
    draw.text(
        (555 * scale, 680 * scale),
        "结构化事件事实",
        font=small_font,
        fill="#8C5F1B",
    )
    draw.text(
        (1278 * scale, 680 * scale),
        "情报卡文本",
        font=small_font,
        fill="#8C5F1B",
    )

    image.save(path, quality=95)


def build_docx() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    QA_DIR.mkdir(parents=True, exist_ok=True)

    timeline = ASSET_DIR / "timeline.png"
    ui = ASSET_DIR / "ui.png"
    agent = ASSET_DIR / "agent.png"
    build_timeline(timeline)
    build_ui(ui)
    build_agent_pipeline(agent)

    doc = Document()
    base.configure_styles(doc)
    base.configure_page(doc)
    add_footer(doc.sections[0].footer.paragraphs[0])

    props = doc.core_properties
    props.title = "芯片战争不完全信息动态博弈最终游戏设计方案"
    props.subject = "游戏规则、交互流程、智能体架构与选举换台机制"
    props.author = ""
    props.keywords = "芯片战争, 不完全信息, 动态博弈, 智能体, 大模型, 游戏设计"

    title = doc.add_paragraph(style="Title")
    title_run = title.add_run("芯片战争不完全信息动态博弈")
    title_run.add_break()
    title_run.add_text("最终游戏设计方案")

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("面向中国玩家与智能体对手的可玩规则设计")

    meta = doc.add_paragraph()
    meta.paragraph_format.space_before = Pt(5)
    meta.paragraph_format.space_after = Pt(15)
    meta_run = meta.add_run("策略游戏版本 1.0  |  2026 年 10 月 28 日")
    meta_run.font.name = BODY_FONT
    meta_run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
    meta_run.font.size = Pt(9.5)
    meta_run.font.color.rgb = MUTED

    doc.add_heading("设计摘要", level=1)
    doc.add_paragraph(
        "玩家控制中国政府，在六轮政策周期中观察受控生成的世界事件、美国 Agent 信号和已发生的政策行动，"
        "推断美国当前属于安全优先型、产业敏感型还是条件交换型，并据此选择产业、管制与外交政策。"
    )
    doc.add_paragraph(
        "第三轮结束后触发美国执政联盟更替。大选结果公开，新政府真实类型仍然隐藏，旧证据保留但权重下降，"
        "美国政策进入一轮过渡期。玩家需要在第四轮重新判断新政府，并调整前三轮遗留的供应链路径。"
    )
    doc.add_paragraph(
        "本方案保留文档 2 的收益结算作为国家绩效层，在其上增加隐藏类型、信念、局部博弈树和 Agent 决策层。"
        "大模型只负责把结构化事件改写成具体情境，不决定隐藏类型、概率和结算结果。"
    )

    doc.add_heading("五项最终决策", level=2)
    base.reset_numbering()
    base.add_numbered(doc, "采用六轮结构，第三轮结束触发美国执政联盟更替。")
    base.add_numbered(doc, "美国使用安全优先、产业敏感、条件交换三种稳定类型。")
    base.add_numbered(doc, "每轮一条世界事件、一条 Agent 信号、一道预测题和三张政策卡。")
    base.add_numbered(doc, "后端使用局部博弈树和期望效用选择动作，界面只显示关键分支。")
    base.add_numbered(doc, "大模型负责情境叙事，规则引擎负责事实、概率、合法动作和结算。")

    doc.add_heading("文档用途", level=2)
    doc.add_paragraph(
        "本文件用于游戏规则评审、Agent 架构设计、界面原型和首版开发排期。"
        "“分析沙盘模式”继续服务透明规则与对照实验，“策略游戏模式”用于加入不确定性和玩家判断。"
    )

    section_heading(doc, "一  最终定位与运行模式")
    doc.add_paragraph(
        "游戏的核心不是让玩家看完所有模型数据，而是让玩家在信息不完整时作出可解释的政策选择。"
        "玩家需要判断美国当前意图，在产业损失、技术安全和外交空间之间取舍，并承担预测错误与政策路径依赖。"
    )
    base.add_caption(doc, "表 1  两种运行模式", before=True)
    base.add_table(
        doc,
        ["模式", "信息条件", "用途"],
        [
            [
                "分析沙盘",
                "类型固定或公开，规则透明，不展示隐藏偏好。",
                "研究、对照实验、政策教学和参数敏感性分析。",
            ],
            [
                "策略游戏",
                "美国类型隐藏，第三轮后换台，情境受控随机。",
                "玩家判断、Agent 对手、分支预测和沉浸式复盘。",
            ],
        ],
        [1.2, 2.8, 2.85],
        font_size=9.0,
    )
    doc.add_paragraph(
        "两种模式共用政策动作、供应链、企业行为和收益函数。策略游戏模式不能把游戏结果包装为现实政策预测。"
    )

    doc.add_heading("核心规则", level=2)
    base.add_caption(doc, "表 2  游戏核心规则", before=True)
    base.add_table(
        doc,
        ["项目", "设定"],
        [
            ["对局长度", "六轮。前三轮为旧政府，第四轮为过渡政府，第五至六轮为新政府。"],
            ["玩家角色", "中国政府。玩家拥有 2 AP，每轮最多执行两项主政策。"],
            ["主政策卡", "产业、管制、外交三类各一张，其他政策进入扩展抽屉。"],
            ["核心推理", "每轮一道美国行为预测题，预测锁定后才能选择政策。"],
            ["美国类型", "安全优先型、产业敏感型、条件交换型。"],
            ["其他 Agent", "荷兰、日本、韩国政府及美国、中国企业和台积电。"],
        ],
        [1.35, 5.5],
        font_size=9.0,
    )
    doc.add_heading("玩家目标", level=2)
    base.add_bullet(doc, "维持可用算力和国产有效交付，避免先进封装、制造设备或 HBM 成为持续瓶颈。")
    base.add_bullet(doc, "控制产业压力和财政成本，避免为了短期安全目标透支后续政策空间。")
    base.add_bullet(doc, "在前三轮为美国换台保留调整余地，第四轮根据新政府信号重新配置资源。")

    section_heading(doc, "二  六轮结构与选举冲击")
    doc.add_paragraph(
        "大选是公开事件，新政府类型是隐藏状态。玩家提前知道第三轮结束会发生换台，因此前三轮不仅要判断当前政府，"
        "还要为下一届政府保留调整空间。"
    )
    base.add_figure(doc, timeline, 6.45, "图 1  六轮结构与美国执政联盟更替")
    base.add_caption(doc, "表 3  各阶段玩法重点", before=True)
    base.add_table(
        doc,
        ["阶段", "美国状态", "玩家任务"],
        [
            ["第 1 至 3 轮", "旧政府隐藏类型稳定", "积累证据，判断旧政府，控制供应链暴露。"],
            ["第 3 轮结束", "执政联盟更替公开", "观察选举结果、新政府纲领和人事信号。"],
            ["第 4 轮", "新政府过渡，政策惯性", "重新判断新类型，决定等待或提前调整。"],
            ["第 5 至 6 轮", "新政府类型逐步稳定", "利用已经形成的路径，争取最终目标。"],
        ],
        [1.35, 2.2, 3.3],
        font_size=9.0,
    )
    doc.add_paragraph(
        "新类型通过条件转移概率决定。旧政府类型、选举结果和第三轮经济状态共同影响新政府分布。"
        "第四轮保留政策惯性，避免一次事件把前三轮的判断和供应链影响全部抹掉。"
    )
    base.add_bullet(doc, "选举结果公开，但新政府真实类型继续隐藏。")
    base.add_bullet(doc, "旧政府证据仍然保留，新政府证据在第四轮获得更高权重。")
    base.add_bullet(doc, "第四轮只部分执行新政府偏好，第五轮起逐步稳定。")

    section_heading(doc, "三  参与主体与隐藏类型")
    base.add_caption(doc, "表 4  参与主体与行动逻辑", before=True)
    base.add_table(
        doc,
        ["主体", "行动逻辑"],
        [
            ["中国玩家", "根据证据判断美国类型，选择政策组合和公开表态。"],
            ["美国政府 Agent", "依据隐藏类型、国内压力、中国信号和历史承诺选择管制与协调政策。"],
            ["荷兰、日本、韩国", "依据商业暴露、美国压力、中国反制和本国产业收益选择配合强度。"],
            ["美国与中国企业", "依据许可、利润、库存、客户和供应商关系调整经营。"],
            ["台积电", "依据产能、客户承诺、海外布局和管制风险分配产能。"],
        ],
        [1.55, 5.3],
        font_size=9.0,
    )
    base.add_caption(doc, "表 5  美国隐藏类型", before=True)
    base.add_table(
        doc,
        ["类型", "核心偏好", "可观察行为"],
        [
            ["安全优先型", "技术安全、盟友协调、长期限制", "扩大许可审查，增加执法预算，对产业损失容忍度高。"],
            ["产业敏感型", "企业利润、就业、供应链稳定", "保留豁免，延迟执行，对国内产业压力反应强。"],
            ["条件交换型", "对等让步、盟友支持、政策交换", "根据中国让步和盟友协调改变管制强度。"],
        ],
        [1.25, 2.2, 3.4],
        font_size=9.0,
    )
    doc.add_paragraph(
        "类型决定 Agent 的效用权重和动作分布，不能只改变政策解释。类型在每届政府内保持稳定，"
        "第三轮结束时才允许发生一次有条件转换。"
    )
    doc.add_paragraph(
        "玩家看到的是类型概率、证据来源和历史行为，不看到真实类型和内部权重。"
        "盟国政府和企业使用规则 Agent，第一版不增加它们的隐藏类型。"
    )

    section_heading(doc, "四  单轮互动流程")
    base.add_caption(doc, "表 6  单轮玩家操作流程", before=True)
    base.add_table(
        doc,
        ["步骤", "界面行为", "玩家操作"],
        [
            ["1  事件出现", "展示世界事件、来源、可信度和证据标签。", "阅读证据，不需要立即决策。"],
            ["2  类型研判", "展示当前美国类型概率和证据板。", "把证据归入支持或反对某类型的假设。"],
            ["3  锁定判断", "显示一道核心预测题。", "给出概率并锁定，之后不能修改。"],
            ["4  政策选择", "展示三张主卡、AP、预算和分支结果。", "选择 1 至 2 张卡，形成政策组合。"],
            ["5  公开表态", "强硬、模糊、合作三种措辞。", "选择是否释放信号，并承担可信度变化。"],
            ["6  提交", "确认预测、政策和表态。", "提交后进入封闭结算。"],
            ["7  行动揭示", "揭示美国政策、盟国反应和企业调整。", "观察，不进行操作。"],
            ["8  因果战报", "显示前三影响、预测对比和信念更新。", "判断本轮推理是否成立。"],
            ["9  下轮准备", "显示下轮预警和阶段目标。", "决定继续坚持还是调整战略。"],
        ],
        [1.3, 3.1, 2.45],
        font_size=8.7,
    )
    doc.add_paragraph(
        "判断必须先于政策选择。这样玩家会先思考美国可能的行动，再判断哪张政策卡最能承受不同类型分支，"
        "而不是根据即时收益直接挑选最强政策。"
    )
    doc.add_heading("判断力计分", level=2)
    doc.add_paragraph(
        "玩家的预测与实际结果单独形成判断力评分，不计入国家绩效。概率越接近实际结果，评分越高；"
        "如果预测错误但政策在另一分支下仍可承受，复盘会分别评价“判断质量”和“决策质量”。"
    )
    base.add_bullet(doc, "每轮记录玩家给出的概率、美国实际行为和系统后验。")
    base.add_bullet(doc, "结算后显示预测误差，并保留最近三轮的判断趋势。")
    base.add_bullet(doc, "判断力评分只用于复盘和教学，不替代国家绩效。")

    section_heading(doc, "五  决策界面与证据板")
    doc.add_paragraph(
        "主界面保持一屏决策。左侧供应链地图承担局势读取，右侧情报板承担不确定性，底部政策卡承担行动选择。"
        "玩家不需要在多个页面之间来回寻找信息。"
    )
    base.add_figure(doc, ui, 6.45, "图 2  新政府过渡期的决策主界面示意")
    doc.add_heading("界面区域", level=2)
    base.add_bullet(doc, "顶部：轮次、阶段任务、三项角色指标和当前政权状态。")
    base.add_bullet(doc, "左侧：供应链地图、瓶颈节点和三类镜头。")
    base.add_bullet(doc, "右侧：公开信号、证据标签、类型概率和本轮预测。")
    base.add_bullet(doc, "底部：三张政策卡、AP、预算、公开表态和提交按钮。")
    doc.add_paragraph(
        "大选后旧证据保留在证据板上并标记为“旧政府时期”，新证据优先更新当前判断。"
        "玩家可以展开完整证据历史，但主屏只显示最近三条。"
    )

    section_heading(doc, "六  受控情境生成")
    doc.add_paragraph(
        "情境可以随机生成，但必须受规则约束。随机的是本轮出现哪类证据，"
        "不是大模型凭空发明一个与状态无关的新事实。"
    )
    base.add_caption(doc, "表 7  三类决策依据", before=True)
    base.add_table(
        doc,
        ["依据", "示例", "对判断的作用"],
        [
            ["世界事件", "出口数据异常、设备订单积压、企业盈利变化", "提供产业状态线索，不直接说明美国类型。"],
            ["Agent 信号", "官员讲话、盟友磋商、匿名泄密、政策吹风", "可能提供意图线索，也可能是虚张声势。"],
            ["已发生行动", "扩大许可审查、给予豁免、增加执法预算", "最强行为证据，结算后更新类型概率。"],
        ],
        [1.35, 3.2, 2.3],
        font_size=9.0,
    )
    doc.add_heading("生成流程", level=2)
    base.add_bullet(doc, "规则引擎根据当前状态、美国类型和事件权重选取结构化事件模板。")
    base.add_bullet(doc, "模板携带来源类型、可信度、可见事实、证据标签和似然参数。")
    base.add_bullet(doc, "大模型根据模板生成新闻、讲话、企业公告或简报文本。")
    base.add_bullet(doc, "校验器检查文本没有泄露真实类型，也没有新增数字、事实和行为主体。")
    base.add_bullet(doc, "界面把结构化事实和生成文本合成为情报卡，交给玩家判断。")
    base.add_bullet(doc, "结算后由规则引擎使用似然参数更新类型概率，不由大模型输出概率。")
    doc.add_paragraph(
        "每局使用固定随机种子。相同种子、相同政策和相同类型应产生可复盘的情境序列，"
        "以便调试、平衡和比较不同玩家策略。"
    )
    doc.add_heading("示例情报卡", level=2)
    base.add_lead_paragraph(
        doc,
        "结构化事实：",
        "美国推动日本和荷兰统一敏感设备许可口径；来源为盟友政策简报；可信度 70%；"
        "证据标签为盟友协调、管制升级；影响美国设备许可和中国材料许可。",
    )
    base.add_lead_paragraph(
        doc,
        "生成文本：",
        "美国正要求日本和荷兰统一敏感设备许可口径。多名参与磋商的人士称，"
        "美方希望在下轮政策调整前完成协调。",
    )

    section_heading(doc, "七  局部博弈树与 Agent 架构")
    doc.add_paragraph(
        "完整六轮五方博弈树不适合在浏览器中实时求解。游戏使用局部博弈树，"
        "只展开当前政策、美国类型和第三方主要反应。"
    )
    base.add_figure(doc, agent, 6.45, "图 3  Agent 决策与局部博弈树流程")
    base.add_caption(doc, "表 8  博弈树的两种用途", before=True)
    base.add_table(
        doc,
        ["层级", "内容", "使用方式"],
        [
            ["实时局部树", "中国候选政策 × 美国类型 × 第三方反应", "生成分支结果、最低结果和期望结果。"],
            ["离线完整模型", "六轮状态、历史、类型转换和全部候选动作", "PBE、CFR、逆向归纳和参数平衡测试。"],
            ["玩家界面", "两到三条主要分支", "帮助玩家理解不确定性，不展示内部求解细节。"],
        ],
        [1.25, 3.05, 2.55],
        font_size=8.9,
    )
    doc.add_paragraph(
        "每个 Agent 至少包含私有类型、当前信念、记忆、合法动作集和效用权重。"
        "美国 Agent 依据当前类型和公开信号选择政策，盟国与企业 Agent 依据政策结果和自身状态反应。"
    )
    doc.add_paragraph(
        "大选换台后，美国 Agent 的类型和效用权重更新，但政策承诺、既有管制和企业预期保留。"
        "第四轮只部分执行新政府偏好，形成一轮可观察的政策惯性。"
    )

    section_heading(doc, "八  目标函数与 Agent 效用")
    doc.add_paragraph(
        "文档 2 的收益函数保留为世界结算层。三种角色分别评价联合限制收益、商业损失、"
        "资源损失、施压成本、配合收益和自身反制成本。六轮总分用于结局评价与情景比较。"
    )
    base.add_caption(doc, "表 9  国家结算层", before=True)
    base.add_table(
        doc,
        ["角色", "收益构成"],
        [
            ["美国", "联合限制战略收益 - 商业暴露成本 - 资源成本 - 施压执行成本"],
            ["荷兰、日本、韩国", "配合收益 - 商业暴露成本 - 资源成本 - 不配合成本"],
            ["中国", "负联合限制强度 - 自身资源反制成本"],
        ],
        [1.5, 5.35],
        font_size=9.0,
    )
    doc.add_paragraph(
        "国家结算分不能直接作为 Agent 的唯一奖励。中国的结算收益中，反制只增加自身成本，"
        "完全理性的最后一轮会选择零反制。因此还需增加独立的 Agent 决策效用。"
    )
    base.add_caption(doc, "表 10  国家绩效与 Agent 效用的分工", before=True)
    base.add_table(
        doc,
        ["概念", "用途", "是否公开"],
        [
            ["国家绩效", "评价一局结果，生成总分、分项和结局等级。", "结算后公开。"],
            ["Agent 效用", "驱动隐藏类型的行为选择，可包含安全、产业、外交、声誉等权重。", "不公开具体权重。"],
        ],
        [1.35, 4.25, 1.25],
        font_size=9.0,
    )
    base.add_caption(doc, "表 11  类型权重的主要差异", before=True)
    base.add_table(
        doc,
        ["类型", "高权重目标", "低权重目标"],
        [
            ["安全优先型", "技术安全、联盟协调", "企业短期利润"],
            ["产业敏感型", "企业收入、就业、供应链稳定", "长期技术限制收益"],
            ["条件交换型", "对等让步、政策交换、可信承诺", "单边升级收益"],
        ],
        [1.35, 3.0, 2.5],
        font_size=9.0,
    )
    doc.add_paragraph(
        "大模型可以解释政策，也可以把类型权重转化为政策讲话。隐藏类型、概率、动作选择和收益结算必须由规则引擎完成。"
    )

    section_heading(doc, "九  结局评价与首版范围")
    base.add_caption(doc, "表 12  结局评价维度", before=True)
    base.add_table(
        doc,
        ["维度", "评价内容"],
        [
            ["国家绩效", "可用算力、国产交付、产业压力、财政和供应链风险。"],
            ["判断质量", "预测概率与实际结果的误差，以及证据使用的合理性。"],
            ["适应能力", "大选后重新判断新政府并调整政策的速度与效果。"],
            ["政策路径", "前三轮选择造成的库存、供应链、盟友关系与承诺。"],
        ],
        [1.35, 5.5],
        font_size=9.0,
    )
    doc.add_heading("首版交付范围", level=2)
    base.add_bullet(doc, "中国玩家，六轮对局，第三轮结束触发美国执政联盟更替。")
    base.add_bullet(doc, "美国三种隐藏类型，第四轮政策惯性，第五至六轮新政府定型。")
    base.add_bullet(doc, "每轮一条世界事件、一条 Agent 信号、一道预测题和三张政策卡。")
    base.add_bullet(doc, "局部博弈树展示两到三条政策分支，不显示完整树。")
    base.add_bullet(doc, "盟国与企业使用规则 Agent，暂不增加隐藏类型。")
    base.add_bullet(doc, "大模型只生成情境文本，并保留规则回退和动作校验。")

    doc.add_heading("验收标准", level=2)
    base.add_bullet(doc, "玩家每轮能够说出美国当前可能是哪一型，以及依据是什么。")
    base.add_bullet(doc, "玩家在大选后能够重新形成判断，而不是继续沿用旧类型概率。")
    base.add_bullet(doc, "每项情报都能追溯到来源、可信度和对应证据标签。")
    base.add_bullet(doc, "每张政策卡都能展示不同美国类型下的结果区间。")
    base.add_bullet(doc, "所有大模型文本均通过状态、事实和动作合法性校验。")
    base.add_bullet(doc, "结算后同时显示玩家预测、系统后验和实际行为。")

    doc.add_heading("依据与边界", level=2)
    doc.add_paragraph(
        "本方案沿用《五国半导体出口管制互动沙盘设计文件》的收益结构，并参考五智能体 PBE 示例的信息集和延续博弈写法。"
        "策略游戏模式是产品扩展，不用于声称现实国家行为、真实均衡或经济损失预测。"
    )

    doc.save(DOCX_PATH)


if __name__ == "__main__":
    build_docx()
    print(DOCX_PATH)
